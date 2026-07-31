"""Convert supported source files into structured parsed JSON."""

from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import tempfile
import time
import uuid

from datetime import datetime, timezone
from html.parser import HTMLParser
from multiprocessing import get_context
from pathlib import Path

# Keep stdlib typing loaded before Docling; OpenCV's loader temporarily
# exposes cv2/typing as a top-level import path in fresh spawn workers.
from typing import Any

import docling
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    PdfPipelineOptions,
    TableStructureOptions,
)
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling_core.types.doc import DocItemLabel
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from rich.console import Console
import xlrd

from pipeline.config import (
    CRAWLER_DATA_DIR,
    CRAWLER_DB_PATH,
    PARSED_DATA_DIR,
    PIPELINE_DATA_DIR,
    PARSE_SUPPORTED_FILE_FORMATS,
    SOURCE_TYPE_DIRS,
)
from pipeline.crawler_registry import normalize_local_path, read_registry_rows
from pipeline.parse.post_process import extract_section_number
from pipeline.parse.worker import parse_supported_file_worker
from pipeline.corpus_manifest import SOURCE_FAMILY_MAP

console = Console()
PARSE_CACHE_VERSION = 5
MIGRATABLE_PARSE_CACHE_VERSIONS = frozenset({4})
DEFAULT_PARSE_TIMEOUT_SECONDS = 600
WORKER_ERROR_MESSAGE_LIMIT = 4096
WORKER_CLEANUP_RESERVE_SECONDS = 5
WORKER_TERMINATE_GRACE_SECONDS = 2
WORKER_NORMAL_EXIT_GRACE_SECONDS = 1
MIN_SUBSTANTIVE_PAGE_CHARS = 40
ACTIVE_REGISTRY_STATUSES = frozenset({"downloaded", "updated", "unchanged", "error"})


def create_converter(*, do_ocr: bool = False) -> DocumentConverter:
    """Create a Docling converter; callers opt into expensive OCR explicitly."""
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = do_ocr
    pipeline_options.do_table_structure = True
    pipeline_options.table_structure_options = TableStructureOptions(
        do_cell_matching=True,
    )
    return DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options),
        }
    )


def file_sha256(path: Path) -> str:
    """Compute SHA-256 hex digest of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def artifact_stem_for_source_path(source_path: Path) -> str:
    """Build a parsed artifact stem that is stable across supported formats."""
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        return source_path.stem
    return f"{source_path.stem}__{suffix.lstrip('.')}"


def parsed_output_paths(source_path: Path, source_type: str) -> tuple[Path, Path, Path]:
    out_dir = PARSED_DATA_DIR / SOURCE_TYPE_DIRS.get(source_type, source_type)
    out_dir.mkdir(parents=True, exist_ok=True)

    stem = artifact_stem_for_source_path(source_path)
    return (
        out_dir / f"{stem}.json",
        out_dir / f"{stem}.md",
        out_dir / f"{stem}.meta.json",
    )


def get_parse_timeout_seconds() -> int:
    raw_value = os.getenv("PIPELINE_PARSE_TIMEOUT_SECONDS")
    if raw_value is None:
        return DEFAULT_PARSE_TIMEOUT_SECONDS

    try:
        return max(0, int(raw_value))
    except ValueError:
        return DEFAULT_PARSE_TIMEOUT_SECONDS


def should_parse(source_path: Path, meta_path: Path) -> bool:
    """Check if a source file needs (re-)parsing by comparing file hashes."""
    if not meta_path.exists():
        return True
    try:
        meta = json.loads(meta_path.read_text())
        if meta.get("parse_status") != "parsed" or not meta.get("answer_ready", False):
            return True
        stored_hash = meta.get("source_file_hash") or meta.get("source_pdf_hash") or ""
        current_hash = f"sha256:{file_sha256(source_path)}"
        if stored_hash != current_hash:
            return True
        json_path, md_path, _ = parsed_output_paths(
            source_path, source_type_for_file(source_path)
        )
        if meta.get("parse_cache_version") != PARSE_CACHE_VERSION:
            return not _migrate_legacy_parsed_sidecar(
                meta_path=meta_path,
                meta=meta,
                json_path=json_path,
                md_path=md_path,
            )
        artifacts = meta.get("artifacts")
        if not isinstance(artifacts, dict):
            return True
        for name, path in (("json", json_path), ("markdown", md_path)):
            expected = artifacts.get(name, {}).get("sha256")
            if not isinstance(expected, str) or not path.is_file():
                return True
            if f"sha256:{file_sha256(path)}" != expected:
                return True
        try:
            payload = json.loads(json_path.read_text())
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            return True
        return not isinstance(payload, dict)
    except (OSError, json.JSONDecodeError, KeyError, TypeError):
        return True


def _new_synthetic_section(title: str | None, level: int) -> dict:
    normalized_title = title.strip() if title else None
    return {
        "section_number": extract_section_number(normalized_title)
        if normalized_title
        else None,
        "title": normalized_title,
        "level": level,
        "content_text": "",
        "has_tables": False,
        "has_definitions": bool(
            normalized_title and "definition" in normalized_title.lower()
        ),
    }


def _append_section_content(section: dict, text: str) -> None:
    if not text.strip():
        return
    prefix = "\n" if section["content_text"] else ""
    section["content_text"] += f"{prefix}{text.strip()}"


def _render_docling_table(table_entry: dict | None) -> str:
    if not isinstance(table_entry, dict):
        return ""

    table_cells = ((table_entry.get("data") or {}).get("table_cells")) or []
    if not isinstance(table_cells, list):
        return ""

    rows: dict[int, dict[int, str]] = {}
    max_col = -1
    for cell in table_cells:
        if not isinstance(cell, dict):
            continue
        row_idx = cell.get("start_row_offset_idx")
        col_idx = cell.get("start_col_offset_idx")
        if not isinstance(row_idx, int) or not isinstance(col_idx, int):
            continue

        cell_text = " ".join(str(cell.get("text") or "").split())
        if not cell_text:
            continue

        rows.setdefault(row_idx, {})[col_idx] = cell_text
        end_col = cell.get("end_col_offset_idx")
        if isinstance(end_col, int):
            max_col = max(max_col, end_col - 1)
        else:
            max_col = max(max_col, col_idx)

    if not rows or max_col < 0:
        return ""

    rendered_rows: list[str] = []
    for row_idx in sorted(rows):
        row = rows[row_idx]
        columns = [row.get(col_idx, "") for col_idx in range(max_col + 1)]
        if any(column.strip() for column in columns):
            rendered_rows.append(" | ".join(columns).strip())

    return "\n".join(rendered_rows).strip()


def _compute_quality_flags(
    sections: list[dict],
    *,
    reported_section_count: int | None = None,
) -> list[str]:
    quality_flags: list[str] = []
    section_count = (
        reported_section_count if reported_section_count is not None else len(sections)
    )
    if section_count == 0:
        quality_flags.append("zero_section_parse")

    content_char_count = sum(
        len((section.get("content_text") or "").strip()) for section in sections
    )
    if content_char_count == 0:
        quality_flags.append("title_only_or_blank_body")

    return quality_flags


def _render_synthetic_markdown(sections: list[dict]) -> str:
    parts: list[str] = []
    for index, section in enumerate(sections):
        title = section.get("title")
        content = (section.get("content_text") or "").strip()
        level = max(1, min(int(section.get("level", 1) or 1), 6))

        if title:
            block = f"{'#' * level} {title}"
            if content:
                block += f"\n\n{content}"
            parts.append(block)
        elif content:
            parts.append(content)
        else:
            parts.append(f"Section {index + 1}")

    return "\n\n".join(part for part in parts if part).strip()


def _parse_pdf_once(pdf_path: Path, converter: DocumentConverter) -> dict:
    """Parse a single PDF and return the parsed payload plus metadata."""
    start = time.time()
    result = converter.convert(str(pdf_path))
    elapsed = time.time() - start

    doc = result.document
    doc_dict = doc.export_to_dict()
    md_text = doc.export_to_markdown()
    table_entries = (
        iter(doc_dict.get("tables", []))
        if isinstance(doc_dict.get("tables"), list)
        else iter(())
    )

    section_count = 0
    table_count = 0
    content_char_count = 0
    synthetic_sections: list[dict] = []
    current_section: dict | None = None
    page_content: dict[int, list[str]] = {}

    for item, level in doc.iterate_items():
        label = getattr(item, "label", None)
        text = (getattr(item, "text", "") or "").strip()
        provenance = getattr(item, "prov", None) or []
        item_pages = {
            int(getattr(prov, "page_no", 0))
            for prov in provenance
            if isinstance(getattr(prov, "page_no", None), int)
            and int(getattr(prov, "page_no")) > 0
        }

        if label in (DocItemLabel.SECTION_HEADER, DocItemLabel.TITLE):
            if current_section is not None:
                synthetic_sections.append(current_section)
            current_section = _new_synthetic_section(text or None, level or 1)
            section_count += 1
        elif label == DocItemLabel.TABLE:
            table_count += 1
            if current_section is None:
                current_section = _new_synthetic_section(None, 0)
            current_section["has_tables"] = True
            table_text = _render_docling_table(next(table_entries, None))
            if table_text:
                _append_section_content(current_section, table_text)
                content_char_count += len(table_text)
                for page_no in item_pages:
                    page_content.setdefault(page_no, []).append(table_text)
        elif label in (
            DocItemLabel.PARAGRAPH,
            DocItemLabel.LIST_ITEM,
            DocItemLabel.TEXT,
            DocItemLabel.FOOTNOTE,
            DocItemLabel.CAPTION,
            DocItemLabel.REFERENCE,
            DocItemLabel.CODE,
            DocItemLabel.FORMULA,
        ):
            if current_section is None:
                current_section = _new_synthetic_section(None, 0)
            _append_section_content(current_section, text)
            content_char_count += len(text)
            for page_no in item_pages:
                page_content.setdefault(page_no, []).append(text)

    if current_section is not None:
        synthetic_sections.append(current_section)

    _num_pages = getattr(doc, "num_pages", None)
    if callable(_num_pages):
        page_count = _num_pages()
    elif _num_pages is not None:
        page_count = _num_pages
    else:
        page_count = len(getattr(doc, "pages", []))

    reported_section_count = (
        section_count if section_count > 0 else len(synthetic_sections)
    )
    quality_flags = _compute_quality_flags(
        synthetic_sections,
        reported_section_count=reported_section_count,
    )

    return {
        "payload": doc_dict,
        "markdown": md_text,
        "section_count": reported_section_count,
        "table_count": table_count,
        "page_count": page_count,
        "content_char_count": content_char_count,
        "parse_time_seconds": round(elapsed, 2),
        "source_format": "pdf",
        "parse_backend": "docling",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
        "page_content": {
            page_no: " ".join(parts) for page_no, parts in page_content.items()
        },
    }


def _normalize_pdf_text(text: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", text.lower()))


def _page_has_images(page: Any) -> bool:
    try:
        resources = page.get("/Resources") or {}
        xobjects = resources.get("/XObject") or {}
        xobjects = (
            xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
        )
        return any(
            (obj.get_object() if hasattr(obj, "get_object") else obj).get("/Subtype")
            == "/Image"
            for obj in xobjects.values()
        )
    except Exception:
        return True


def evaluate_no_ocr_pdf(pdf_path: Path, parsed: dict) -> tuple[bool, str, list[dict]]:
    """Prove every nonblank source page survived no-OCR parsing, or fail closed."""
    evidence: list[dict] = []
    try:
        pages = PdfReader(str(pdf_path), strict=False).pages
        for page_no, page in enumerate(pages, start=1):
            source = _normalize_pdf_text(page.extract_text() or "")
            extracted = _normalize_pdf_text(
                parsed.get("page_content", {}).get(page_no, "")
            )
            has_images = _page_has_images(page)
            source_tokens = source.split()
            extracted_tokens = extracted.split()
            sequence_matches = source_tokens == extracted_tokens
            coverage = 1.0 if sequence_matches else 0.0
            source_chars = len(source.replace(" ", ""))
            extracted_chars = len(extracted.replace(" ", ""))
            blank = source_chars == 0 and not has_images
            evidence.append(
                {
                    "page": page_no,
                    "source_text_chars": source_chars,
                    "docling_text_chars": extracted_chars,
                    "normalized_token_coverage": round(coverage, 3),
                    "normalized_token_sequence_matches": sequence_matches,
                    "has_images": has_images,
                    "source_blank": blank,
                }
            )
            if blank:
                continue
            if source_chars < MIN_SUBSTANTIVE_PAGE_CHARS:
                return False, f"page_{page_no}_image_or_sparse_source", evidence
            if extracted_chars < MIN_SUBSTANTIVE_PAGE_CHARS:
                return False, f"page_{page_no}_missing_substantive_provenance", evidence
            character_coverage = extracted_chars / source_chars if source_chars else 1.0
            evidence[-1]["normalized_character_coverage"] = round(character_coverage, 3)
            if not sequence_matches or character_coverage != 1.0:
                return False, f"page_{page_no}_insufficient_text_coverage", evidence
            if int(parsed.get("table_count", 0)) > 0:
                return False, "table_structure_requires_ocr", evidence
    except Exception as error:
        evidence.append(
            {"preflight_error": f"{type(error).__name__}: {str(error)[:500]}"}
        )
        return False, "pypdf_preflight_failed", evidence
    if not evidence:
        return False, "pdf_has_no_pages", evidence
    return True, "all_nonblank_pages_match_source_token_sequence_exactly", evidence


def parse_pdf(pdf_path: Path, converter: DocumentConverter) -> dict:
    """Accept no-OCR only after conservative page-level completeness proof."""
    first = _parse_pdf_once(pdf_path, converter)
    no_ocr_accepted, no_ocr_reason, page_coverage = evaluate_no_ocr_pdf(pdf_path, first)
    attempts = [
        {
            "attempt": 1,
            "ocr": False,
            "parse_time_seconds": first["parse_time_seconds"],
            "quality_flags": first["quality_flags"],
        }
    ]
    selected = first
    ocr_attempted = not (first["answer_ready"] and no_ocr_accepted)
    if ocr_attempted:
        fallback = _parse_pdf_once(pdf_path, create_converter(do_ocr=True))
        attempts.append(
            {
                "attempt": 2,
                "ocr": True,
                "parse_time_seconds": fallback["parse_time_seconds"],
                "quality_flags": fallback["quality_flags"],
            }
        )
        selected = fallback

    selected["ocr_attempted"] = ocr_attempted
    selected["ocr_used"] = ocr_attempted
    selected["no_ocr_acceptance_reason"] = no_ocr_reason if no_ocr_accepted else None
    selected["no_ocr_rejection_reason"] = None if no_ocr_accepted else no_ocr_reason
    selected["page_coverage_evidence"] = page_coverage
    selected.pop("page_content", None)
    selected["parse_attempts"] = attempts
    selected["parse_time_seconds"] = round(
        sum(float(attempt["parse_time_seconds"]) for attempt in attempts), 2
    )
    return selected


def _docx_heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    normalized = style_name.strip().lower()
    if not normalized.startswith("heading "):
        return None
    suffix = normalized.removeprefix("heading ").strip()
    if not suffix.isdigit():
        return None
    return max(1, min(int(suffix), 6))


def parse_docx(docx_path: Path) -> dict:
    """Extract headings and paragraphs from a DOCX into synthetic sections."""
    start = time.time()
    document = WordDocument(str(docx_path))
    sections: list[dict] = []
    current_section: dict | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        level = _docx_heading_level(getattr(paragraph.style, "name", None))
        if level is not None:
            if current_section is not None:
                sections.append(current_section)
            current_section = _new_synthetic_section(text, level)
            continue

        if current_section is None:
            current_section = _new_synthetic_section(None, 0)
        _append_section_content(current_section, text)

    for table in document.tables:
        table_lines: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))

        if not table_lines:
            continue

        if current_section is None:
            current_section = _new_synthetic_section(None, 0)
        _append_section_content(current_section, "\n".join(table_lines))
        current_section["has_tables"] = True

    if current_section is not None:
        sections.append(current_section)

    quality_flags = _compute_quality_flags(
        sections, reported_section_count=len(sections)
    )
    content_char_count = sum(
        len((section.get("content_text") or "").strip()) for section in sections
    )

    return {
        "payload": {
            "synthetic_sections": sections,
            "source_format": "docx",
        },
        "markdown": _render_synthetic_markdown(sections),
        "section_count": len(sections),
        "table_count": sum(1 for section in sections if section.get("has_tables")),
        "page_count": 0,
        "content_char_count": content_char_count,
        "parse_time_seconds": round(time.time() - start, 2),
        "source_format": "docx",
        "parse_backend": "python-docx",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
    }


def parse_xlsx(xlsx_path: Path) -> dict:
    """Extract spreadsheet sheets into synthetic sections with preserved table text."""
    start = time.time()
    workbook = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    sections: list[dict] = []

    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip() if value is not None else "" for value in row
                ]
                if any(values):
                    rows.append(values)

            if not rows:
                continue

            table_text = "\n".join(" | ".join(values) for values in rows)
            section = _new_synthetic_section(sheet.title, 1)
            section["content_text"] = table_text
            section["has_tables"] = True
            sections.append(section)
    finally:
        workbook.close()

    quality_flags = _compute_quality_flags(
        sections, reported_section_count=len(sections)
    )
    content_char_count = sum(
        len((section.get("content_text") or "").strip()) for section in sections
    )

    return {
        "payload": {
            "synthetic_sections": sections,
            "source_format": "xlsx",
        },
        "markdown": _render_synthetic_markdown(sections),
        "section_count": len(sections),
        "table_count": len(sections),
        "page_count": len(sections),
        "content_char_count": content_char_count,
        "parse_time_seconds": round(time.time() - start, 2),
        "source_format": "xlsx",
        "parse_backend": "openpyxl",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
    }


def parse_xls(xls_path: Path) -> dict:
    """Extract legacy XLS workbooks directly into synthetic sections."""
    start = time.time()
    workbook = xlrd.open_workbook(str(xls_path), on_demand=True)
    sections: list[dict] = []
    try:
        for sheet in workbook.sheets():
            rows: list[list[str]] = []
            for row_index in range(sheet.nrows):
                values = [
                    str(sheet.cell_value(row_index, column_index)).strip()
                    for column_index in range(sheet.ncols)
                ]
                if any(values):
                    rows.append(values)
            if not rows:
                continue

            section = _new_synthetic_section(sheet.name, 1)
            section["content_text"] = "\n".join(" | ".join(values) for values in rows)
            section["has_tables"] = True
            sections.append(section)
    finally:
        workbook.release_resources()

    quality_flags = _compute_quality_flags(
        sections,
        reported_section_count=len(sections),
    )
    return {
        "payload": {
            "synthetic_sections": sections,
            "source_format": "xls",
        },
        "markdown": _render_synthetic_markdown(sections),
        "section_count": len(sections),
        "table_count": len(sections),
        "page_count": len(sections),
        "content_char_count": sum(len(section["content_text"]) for section in sections),
        "parse_time_seconds": round(time.time() - start, 2),
        "source_format": "xls",
        "parse_backend": "xlrd",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
    }


def parse_pptx(pptx_path: Path) -> dict:
    """Extract slide text and tables into one synthetic section per slide."""
    start = time.time()
    presentation = Presentation(str(pptx_path))
    sections: list[dict] = []
    table_count = 0

    for index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = (
            title_shape.text.strip()
            if title_shape is not None and title_shape.text
            else f"Slide {index}"
        )
        section = _new_synthetic_section(title, 1)

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape is not title_shape:
                _append_section_content(section, shape.text)
            if getattr(shape, "has_table", False):
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in shape.table.rows
                ]
                _append_section_content(section, "\n".join(rows))
                section["has_tables"] = True
                table_count += 1

        if section["content_text"].strip() or title:
            sections.append(section)

    quality_flags = _compute_quality_flags(
        sections,
        reported_section_count=len(sections),
    )
    return {
        "payload": {
            "synthetic_sections": sections,
            "source_format": "pptx",
        },
        "markdown": _render_synthetic_markdown(sections),
        "section_count": len(sections),
        "table_count": table_count,
        "page_count": len(sections),
        "content_char_count": sum(len(section["content_text"]) for section in sections),
        "parse_time_seconds": round(time.time() - start, 2),
        "source_format": "pptx",
        "parse_backend": "python-pptx",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
    }


def parse_doc(doc_path: Path) -> dict:
    """Convert legacy DOC through headless LibreOffice, then parse the DOCX."""
    libreoffice_binary = os.getenv("LIBREOFFICE_BINARY", "soffice").strip()
    if not libreoffice_binary:
        raise RuntimeError("LIBREOFFICE_BINARY must not be empty")

    with tempfile.TemporaryDirectory(prefix="medicaid-doc-") as temp_dir:
        completed = subprocess.run(
            [
                libreoffice_binary,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                temp_dir,
                str(doc_path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        converted = Path(temp_dir) / f"{doc_path.stem}.docx"
        if completed.returncode != 0 or not converted.exists():
            detail = (completed.stderr or completed.stdout).strip()[:500]
            raise RuntimeError(
                f"LibreOffice DOC conversion failed: {detail or 'no output'}"
            )

        parsed = parse_docx(converted)
        parsed["source_format"] = "doc"
        parsed["parse_backend"] = "libreoffice+python-docx"
        parsed["payload"]["source_format"] = "doc"
        return parsed


class _ArticleHTMLParser(HTMLParser):
    """Extract semantic blocks from the page's main article only."""

    _IGNORED_TAGS = {"script", "style", "nav", "aside", "form"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.article_depth = 0
        self.ignored_depth = 0
        self.heading_level: int | None = None
        self.block_depth = 0
        self.buffer: list[str] = []
        self.sections: list[dict] = []
        self.current_section: dict | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        del attrs
        if tag == "article":
            self.article_depth += 1
            return
        if not self.article_depth:
            return
        if tag in self._IGNORED_TAGS:
            self.ignored_depth += 1
            return
        if self.ignored_depth:
            return
        if len(tag) == 2 and tag[0] == "h" and tag[1].isdigit():
            self.heading_level = int(tag[1])
            self.buffer = []
        elif tag in {"p", "li"}:
            if self.block_depth == 0:
                self.buffer = []
            self.block_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag == "article":
            self.article_depth = max(0, self.article_depth - 1)
            return
        if not self.article_depth:
            return
        if tag in self._IGNORED_TAGS and self.ignored_depth:
            self.ignored_depth -= 1
            return
        if self.ignored_depth:
            return
        if self.heading_level is not None and tag == f"h{self.heading_level}":
            title = " ".join("".join(self.buffer).split())
            if self.current_section is not None:
                self.sections.append(self.current_section)
            self.current_section = _new_synthetic_section(
                title or None, self.heading_level
            )
            self.heading_level = None
            self.buffer = []
        elif tag in {"p", "li"} and self.block_depth:
            self.block_depth -= 1
            if self.block_depth == 0:
                text = " ".join("".join(self.buffer).split())
                if text:
                    if self.current_section is None:
                        self.current_section = _new_synthetic_section(None, 0)
                    _append_section_content(self.current_section, text)
                self.buffer = []

    def handle_data(self, data: str) -> None:
        if (
            self.article_depth
            and not self.ignored_depth
            and (self.heading_level is not None or self.block_depth > 0)
        ):
            self.buffer.append(data)

    def finish(self) -> list[dict]:
        if self.current_section is not None:
            self.sections.append(self.current_section)
            self.current_section = None
        return self.sections


_PAGE_MODIFIED_RE = re.compile(
    r"This page was last modified on\s+(\d{2}/\d{2}/\d{4})",
    re.IGNORECASE,
)


def parse_html(html_path: Path) -> dict:
    """Extract headings and beneficiary-facing text from an official HTML page."""
    start = time.time()
    html = html_path.read_text(encoding="utf-8", errors="replace")
    parser = _ArticleHTMLParser()
    parser.feed(html)
    sections = parser.finish()
    quality_flags = _compute_quality_flags(
        sections, reported_section_count=len(sections)
    )
    content_char_count = sum(
        len((section.get("content_text") or "").strip()) for section in sections
    )
    modified_match = _PAGE_MODIFIED_RE.search(html)
    source_modified_date = None
    if modified_match:
        month, day, year = modified_match.group(1).split("/")
        source_modified_date = f"{year}-{month}-{day}"

    return {
        "payload": {
            "synthetic_sections": sections,
            "source_format": "html",
        },
        "markdown": _render_synthetic_markdown(sections),
        "section_count": len(sections),
        "table_count": 0,
        "page_count": 1,
        "content_char_count": content_char_count,
        "parse_time_seconds": round(time.time() - start, 2),
        "source_format": "html",
        "parse_backend": "stdlib-html-parser",
        "parse_status": "parsed",
        "quality_flags": quality_flags,
        "answer_ready": len(quality_flags) == 0,
        "source_modified_date": source_modified_date,
    }


def _attempt_counts(meta_path: Path, source_hash: str) -> tuple[int, int]:
    try:
        current = json.loads(meta_path.read_text()) if meta_path.exists() else {}
        total = (
            max(
                0,
                int(
                    current.get("attempt_count_total", current.get("attempt_count", 0))
                ),
            )
            + 1
        )
        same_generation = (
            current.get("source_file_hash") == source_hash
            and current.get("parse_cache_version") == PARSE_CACHE_VERSION
        )
        previous_generation = current.get(
            "attempt_count_generation", current.get("attempt_count", 0)
        )
        generation = max(0, int(previous_generation)) + 1 if same_generation else 1
        return generation, total
    except (json.JSONDecodeError, TypeError, ValueError):
        return 1, 1


def _atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _migrate_legacy_parsed_sidecar(
    *,
    meta_path: Path,
    meta: dict[str, Any],
    json_path: Path,
    md_path: Path,
) -> bool:
    """Upgrade a valid v4 success in place without re-running expensive OCR."""
    if meta.get("parse_cache_version") not in MIGRATABLE_PARSE_CACHE_VERSIONS:
        return False
    if (
        not meta.get("answer_ready", False)
        or not json_path.is_file()
        or not md_path.is_file()
    ):
        return False
    try:
        payload = json.loads(json_path.read_text())
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return False
    if not isinstance(payload, dict):
        return False

    migrated = dict(meta)
    migrated.update(
        {
            "parse_cache_version": PARSE_CACHE_VERSION,
            "attempt_count_generation": 0,
            "attempt_count_total": max(0, int(meta.get("attempt_count", 0))),
            "artifacts": {
                "json": {"sha256": f"sha256:{file_sha256(json_path)}"},
                "markdown": {"sha256": f"sha256:{file_sha256(md_path)}"},
            },
            "cache_migrated_at": datetime.now(timezone.utc).isoformat(),
            "cache_migrated_from_version": meta.get("parse_cache_version"),
        }
    )
    _atomic_write_text(meta_path, json.dumps(migrated, indent=2))
    return True


def save_parsed(
    source_path: Path,
    parsed: dict,
    source_type: str,
) -> Path:
    """Save parsed JSON, markdown, and metadata sidecar."""
    json_path, md_path, meta_path = parsed_output_paths(source_path, source_type)

    source_hash = f"sha256:{file_sha256(source_path)}"
    attempt_generation, attempt_total = _attempt_counts(meta_path, source_hash)
    _atomic_write_text(json_path, json.dumps(parsed["payload"], indent=2, default=str))
    _atomic_write_text(md_path, parsed["markdown"])

    meta = {
        "source_file": str(source_path),
        "source_file_hash": source_hash,
        "source_pdf": str(source_path),
        "source_pdf_hash": source_hash,
        "parsed_at": datetime.now(timezone.utc).isoformat(),
        "docling_version": getattr(docling, "__version__", "unknown"),
        "parse_cache_version": PARSE_CACHE_VERSION,
        "attempt_count": attempt_generation,
        "attempt_count_generation": attempt_generation,
        "attempt_count_total": attempt_total,
        "source_format": parsed["source_format"],
        "parse_backend": parsed["parse_backend"],
        "parse_status": parsed["parse_status"],
        "quality_flags": parsed["quality_flags"],
        "answer_ready": parsed["answer_ready"],
        "section_count": parsed["section_count"],
        "table_count": parsed["table_count"],
        "page_count": parsed["page_count"],
        "content_char_count": parsed["content_char_count"],
        "parse_time_seconds": parsed["parse_time_seconds"],
        "source_modified_date": parsed.get("source_modified_date"),
        "warnings": parsed["quality_flags"],
        "ocr_attempted": parsed.get("ocr_attempted", False),
        "ocr_used": parsed.get("ocr_used", False),
        "parse_attempts": parsed.get("parse_attempts", []),
        "no_ocr_acceptance_reason": parsed.get("no_ocr_acceptance_reason"),
        "no_ocr_rejection_reason": parsed.get("no_ocr_rejection_reason"),
        "page_coverage_evidence": parsed.get("page_coverage_evidence", []),
        "artifacts": {
            "json": {"sha256": f"sha256:{file_sha256(json_path)}"},
            "markdown": {"sha256": f"sha256:{file_sha256(md_path)}"},
        },
    }
    _atomic_write_text(meta_path, json.dumps(meta, indent=2))

    return json_path


def save_parse_error(source_path: Path, source_type: str, error: Exception) -> None:
    """Persist parse failure metadata so audit-corpus can classify exceptions explicitly."""
    _, _, meta_path = parsed_output_paths(source_path, source_type)
    source_hash = f"sha256:{file_sha256(source_path)}"
    attempt_generation, attempt_total = _attempt_counts(meta_path, source_hash)
    meta = {
        "source_file": str(source_path),
        "source_file_hash": source_hash,
        "source_pdf": str(source_path),
        "source_pdf_hash": source_hash,
        "parsed_at": datetime.now(timezone.utc).isoformat(),
        "docling_version": getattr(docling, "__version__", "unknown"),
        "parse_cache_version": PARSE_CACHE_VERSION,
        "failure_kind": "timeout" if isinstance(error, TimeoutError) else "exception",
        "attempt_count": attempt_generation,
        "attempt_count_generation": attempt_generation,
        "attempt_count_total": attempt_total,
        "source_format": source_path.suffix.lower().lstrip("."),
        "parse_backend": {
            ".pdf": "docling",
            ".doc": "libreoffice+python-docx",
            ".docx": "python-docx",
            ".xls": "xlrd",
            ".xlsx": "openpyxl",
            ".pptx": "python-pptx",
            ".html": "stdlib-html-parser",
        }.get(source_path.suffix.lower(), "unknown"),
        "parse_status": "parse_exception",
        "quality_flags": ["parse_exception"],
        "answer_ready": False,
        "error_message": str(error),
        "section_count": 0,
        "table_count": 0,
        "page_count": 0,
        "content_char_count": 0,
        "parse_time_seconds": 0,
        "warnings": ["parse_exception"],
    }
    _atomic_write_text(meta_path, json.dumps(meta, indent=2))


def save_unsupported_format(
    source_path: Path,
    source_type: str,
) -> None:
    """Persist an explicit quarantine sidecar for an unsupported artifact."""
    _, _, meta_path = parsed_output_paths(source_path, source_type)
    source_hash = f"sha256:{file_sha256(source_path)}"
    attempt_generation, attempt_total = _attempt_counts(meta_path, source_hash)
    suffix = source_path.suffix.lower() or "(none)"
    meta = {
        "source_file": str(source_path),
        "source_file_hash": source_hash,
        "parsed_at": datetime.now(timezone.utc).isoformat(),
        "parse_cache_version": PARSE_CACHE_VERSION,
        "failure_kind": "unsupported_format",
        "attempt_count": attempt_generation,
        "attempt_count_generation": attempt_generation,
        "attempt_count_total": attempt_total,
        "source_format": suffix.lstrip("."),
        "parse_backend": "unsupported",
        "parse_status": "unsupported_format",
        "quality_flags": ["unsupported_format"],
        "answer_ready": False,
        "error_message": (
            f"Unsupported source format {suffix}; artifact quarantined from loading"
        ),
        "section_count": 0,
        "table_count": 0,
        "page_count": 0,
        "content_char_count": 0,
        "parse_time_seconds": 0,
        "warnings": ["unsupported_format"],
    }
    _atomic_write_text(meta_path, json.dumps(meta, indent=2))


def discover_source_files(source_type: str | None = None) -> list[Path]:
    """Find all supported source files in the crawler's data directory."""
    source_files: list[Path] = []

    def collect_from_directory(dir_name: str) -> None:
        source_dir = CRAWLER_DATA_DIR / dir_name
        if not source_dir.exists():
            return
        for suffix in PARSE_SUPPORTED_FILE_FORMATS:
            source_files.extend(sorted(source_dir.glob(f"*{suffix}")))

    if source_type:
        collect_from_directory(SOURCE_TYPE_DIRS.get(source_type, source_type))
    else:
        for dir_name in SOURCE_TYPE_DIRS.values():
            collect_from_directory(dir_name)

    return source_files


def discover_unsupported_files(
    source_type: str | None = None,
) -> list[Path]:
    """Find downloaded artifacts that have no parser implementation."""
    unsupported: list[Path] = []

    def collect_from_directory(dir_name: str) -> None:
        source_dir = CRAWLER_DATA_DIR / dir_name
        if not source_dir.exists():
            return
        unsupported.extend(
            sorted(
                path
                for path in source_dir.iterdir()
                if path.is_file()
                and not path.name.startswith(".")
                and path.suffix.lower() not in PARSE_SUPPORTED_FILE_FORMATS
            )
        )

    if source_type:
        collect_from_directory(SOURCE_TYPE_DIRS.get(source_type, source_type))
    else:
        for dir_name in SOURCE_TYPE_DIRS.values():
            collect_from_directory(dir_name)

    return unsupported


def source_type_for_file(source_path: Path) -> str:
    """Infer source_type from the source file's parent directory name."""
    dir_name = source_path.parent.name
    for source_type, configured_dir_name in SOURCE_TYPE_DIRS.items():
        if configured_dir_name == dir_name:
            return source_type
    return "unknown"


def parse_supported_file(
    source_path: Path,
    converter: DocumentConverter | None,
) -> dict:
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        if converter is None:
            raise ValueError("PDF parsing requires a DocumentConverter")
        return parse_pdf(source_path, converter)
    if suffix == ".docx":
        return parse_docx(source_path)
    if suffix == ".doc":
        return parse_doc(source_path)
    if suffix == ".xlsx":
        return parse_xlsx(source_path)
    if suffix == ".xls":
        return parse_xls(source_path)
    if suffix == ".pptx":
        return parse_pptx(source_path)
    if suffix == ".html":
        return parse_html(source_path)
    raise ValueError(f"Unsupported parse format: {source_path.suffix}")


def summarize_parsed_output(parsed: dict) -> dict:
    return {
        "section_count": parsed["section_count"],
        "table_count": parsed["table_count"],
        "parse_time_seconds": parsed["parse_time_seconds"],
    }


def parse_and_save_supported_file(
    source_path: Path,
    source_type: str,
    converter: DocumentConverter | None,
) -> dict:
    parsed = parse_supported_file(source_path, converter)
    save_parsed(source_path, parsed, source_type)
    return summarize_parsed_output(parsed)


def _parse_supported_file_worker(
    source_path_str: str,
    source_type: str,
    result_connection: Any,
) -> None:
    source_path = Path(source_path_str)
    try:
        converter = create_converter() if source_path.suffix.lower() == ".pdf" else None
        parsed = parse_supported_file(source_path, converter)
        save_parsed(source_path, parsed, source_type)
        result_connection.send({"ok": True, "summary": summarize_parsed_output(parsed)})
    except Exception as error:  # pragma: no cover - exercised via parent wrapper
        result_connection.send(
            {
                "ok": False,
                "error_message": (f"{type(error).__name__}: {error}")[
                    :WORKER_ERROR_MESSAGE_LIMIT
                ],
            }
        )


def _validate_worker_result(result: Any) -> dict:
    if not isinstance(result, dict) or type(result.get("ok")) is not bool:
        raise RuntimeError("parse worker returned malformed result")
    if result["ok"]:
        summary = result.get("summary")
        if not isinstance(summary, dict):
            raise RuntimeError("parse worker returned malformed success result")
        for key in ("section_count", "table_count", "parse_time_seconds"):
            value = summary.get(key)
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise RuntimeError("parse worker returned malformed success result")
        return result
    message = result.get("error_message")
    if not isinstance(message, str) or not message:
        raise RuntimeError("parse worker returned malformed error result")
    result["error_message"] = message[:WORKER_ERROR_MESSAGE_LIMIT]
    return result


def _worker_named_semaphores(semaphore_prefix: str) -> frozenset[str]:
    """Return named semaphores from one cryptographically unique worker namespace."""
    if os.name != "posix":
        return frozenset()
    file_prefix = f"sem.{semaphore_prefix.lstrip('/')}-"
    try:
        return frozenset(
            f"/{path.name.removeprefix('sem.')}"
            for path in Path("/dev/shm").glob(f"{file_prefix}*")
            if path.is_file()
        )
    except OSError:
        return frozenset()


def _unlink_tracked_semaphore(name: str) -> None:
    """Unlink a killed spawn worker's semaphore and clear its tracker record."""
    from _multiprocessing import sem_unlink
    from multiprocessing import resource_tracker

    try:
        sem_unlink(name)
    except FileNotFoundError:
        pass
    resource_tracker.unregister(name, "semaphore")


def _cleanup_killed_worker_semaphores(semaphore_prefix: str) -> None:
    """Clean named semaphores owned by one reaped, forcibly-stopped worker."""
    for name in sorted(_worker_named_semaphores(semaphore_prefix)):
        try:
            _unlink_tracked_semaphore(name)
        except (OSError, RuntimeError, ValueError) as error:
            console.print(
                f"[yellow]Warning: could not clean worker semaphore {name}: {error}[/yellow]"
            )


def parse_and_save_supported_file_with_timeout(
    source_path: Path,
    source_type: str,
    timeout_seconds: int,
) -> dict:
    ctx = get_context("spawn")
    semaphore_prefix = f"/pipeline-pdf-{os.getpid()}-{uuid.uuid4().hex}"
    parent_connection, child_connection = ctx.Pipe(duplex=False)
    process = ctx.Process(
        target=parse_supported_file_worker,
        args=(str(source_path), source_type, child_connection, semaphore_prefix),
    )
    result: Any = None
    process_started = False
    process_forcibly_stopped = False
    child_connection_closed = False
    deadline = time.monotonic() + max(0, timeout_seconds)
    try:
        process.start()
        process_started = True
        child_connection.close()
        child_connection_closed = True
        while True:
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise TimeoutError(f"parse timed out after {timeout_seconds} seconds")
            if parent_connection.poll(min(0.05, remaining)):
                try:
                    result = parent_connection.recv()
                except (EOFError, OSError) as error:
                    raise RuntimeError(
                        f"parse worker exited without result (exit code {process.exitcode})"
                    ) from error
                break
            if not process.is_alive():
                if parent_connection.poll(0):
                    continue
                raise RuntimeError(
                    f"parse worker exited without result (exit code {process.exitcode})"
                )
    finally:
        if process_started and result is not None and process.is_alive():
            process.join(WORKER_NORMAL_EXIT_GRACE_SECONDS)
        if process_started and process.is_alive():
            process_forcibly_stopped = True
            process.terminate()
            process.join(WORKER_TERMINATE_GRACE_SECONDS)
        if process_started and process.is_alive():
            process.kill()
        if process_started:
            process.join(
                max(0, WORKER_CLEANUP_RESERVE_SECONDS - WORKER_TERMINATE_GRACE_SECONDS)
            )
        worker_reaped = (
            process_started and not process.is_alive() and process.exitcode is not None
        )
        if process_forcibly_stopped and worker_reaped:
            _cleanup_killed_worker_semaphores(semaphore_prefix)
        elif process_forcibly_stopped:
            console.print(
                "[yellow]Warning: parse worker could not be reaped; "
                "leaving its semaphore namespace intact[/yellow]"
            )
        try:
            if not child_connection_closed:
                child_connection.close()
        finally:
            parent_connection.close()

    result = _validate_worker_result(result)
    if not result.get("ok"):
        raise RuntimeError(result.get("error_message") or "parse worker failed")

    return result["summary"]


def _inspect_parse_artifacts(source_path: Path) -> dict[str, Any] | None:
    source_type = source_type_for_file(source_path)
    required = SOURCE_FAMILY_MAP.get(source_type)
    required_for_answering = (
        required.required_for_answering if required is not None else True
    )
    base = {
        "source_file": str(source_path),
        "source_type": source_type,
        "required_for_answering": required_for_answering,
    }
    if not source_path.exists() or not source_path.is_file():
        return {
            **base,
            "parse_status": "missing_source_file",
            "attempt_count_generation": 0,
            "retryable": False,
        }
    json_path, md_path, meta_path = parsed_output_paths(source_path, source_type)
    try:
        meta = json.loads(meta_path.read_text())
    except FileNotFoundError:
        return {
            **base,
            "parse_status": "missing_sidecar",
            "attempt_count_generation": 0,
        }
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {
            **base,
            "parse_status": "malformed_sidecar",
            "attempt_count_generation": 0,
        }
    if not isinstance(meta, dict):
        return {
            **base,
            "parse_status": "malformed_sidecar",
            "attempt_count_generation": 0,
        }
    if (
        meta.get("parse_status") == "parsed"
        and meta.get("answer_ready", False)
        and meta.get("parse_cache_version") != PARSE_CACHE_VERSION
    ):
        # Migrate valid v4 successes before selecting a failure-only recovery
        # batch, so deploying the new cache contract does not reparse thousands
        # of already-good OCR artifacts.
        if not should_parse(source_path, meta_path):
            return None
    attempts = meta.get("attempt_count_generation", meta.get("attempt_count", 0))
    try:
        attempts = max(0, int(attempts))
    except (TypeError, ValueError):
        attempts = 0
    status = meta.get("parse_status")
    if status != "parsed":
        return {
            **base,
            "parse_status": status or "malformed_sidecar",
            "attempt_count_generation": attempts,
        }
    if not meta.get("answer_ready", False):
        return {
            **base,
            "parse_status": "not_answer_ready",
            "attempt_count_generation": attempts,
        }
    current_hash = f"sha256:{file_sha256(source_path)}"
    if (
        meta.get("source_file_hash", meta.get("source_pdf_hash")) != current_hash
        or meta.get("parse_cache_version") != PARSE_CACHE_VERSION
    ):
        return {**base, "parse_status": "stale_sidecar", "attempt_count_generation": 0}
    artifacts = meta.get("artifacts")
    for name, path in (("json", json_path), ("markdown", md_path)):
        if not path.is_file():
            return {
                **base,
                "parse_status": "missing_output_artifact",
                "artifact": name,
                "attempt_count_generation": attempts,
            }
        expected = (
            artifacts.get(name, {}).get("sha256")
            if isinstance(artifacts, dict)
            else None
        )
        if not isinstance(expected, str) or expected != f"sha256:{file_sha256(path)}":
            return {
                **base,
                "parse_status": "invalid_output_artifact",
                "artifact": name,
                "attempt_count_generation": attempts,
            }
    try:
        if not isinstance(json.loads(json_path.read_text()), dict):
            raise ValueError("parsed JSON is not an object")
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        return {
            **base,
            "parse_status": "invalid_output_artifact",
            "artifact": "json",
            "attempt_count_generation": attempts,
        }
    return None


def _active_registry_files(
    source_files: list[Path] | None = None,
    *,
    selected_file: Path | None = None,
    selected_source_type: str | None = None,
) -> list[Path]:
    files = discover_source_files() if source_files is None else source_files
    registry_rows = read_registry_rows()
    if registry_rows:
        active_paths = {
            (
                Path(normalized)
                if Path(normalized).is_absolute()
                else CRAWLER_DATA_DIR / normalized
            )
            for row in registry_rows
            if row.get("status") in ACTIVE_REGISTRY_STATUSES
            and (normalized := normalize_local_path(row.get("local_path")))
        }
        if selected_file is not None:
            files = [selected_file]
        elif selected_source_type is not None:
            files = [
                path
                for path in active_paths
                if source_type_for_file(path) == selected_source_type
            ]
        else:
            files = list(active_paths)
    return sorted(files, key=str)


def inventory_parse_failures(
    source_files: list[Path] | None = None,
    *,
    selected_file: Path | None = None,
    selected_source_type: str | None = None,
) -> list[dict[str, Any]]:
    """Inventory active registry sources, classifying required and optional failures."""
    if CRAWLER_DB_PATH.exists() and not read_registry_rows():
        return [
            {
                "source_file": str(CRAWLER_DB_PATH),
                "source_type": "crawler_registry",
                "parse_status": "empty_registry",
                "required_for_answering": True,
                "attempt_count_generation": 0,
                "retryable": False,
            }
        ]
    return [
        failure
        for path in _active_registry_files(
            source_files,
            selected_file=selected_file,
            selected_source_type=selected_source_type,
        )
        if (failure := _inspect_parse_artifacts(path)) is not None
    ]


def _existing_file_hash(path: Path) -> str | None:
    if not path.exists() or not path.is_file():
        return None
    return f"sha256:{file_sha256(path)}"


def build_parse_inventory_evidence(
    source_files: list[Path] | None = None,
    *,
    selected_file: Path | None = None,
    selected_source_type: str | None = None,
) -> dict[str, Any]:
    """Return deterministic active-source evidence for release mutation checks."""
    if CRAWLER_DB_PATH.exists() and not read_registry_rows():
        records = [
            {
                "source_file": normalize_local_path(str(CRAWLER_DB_PATH)),
                "source_type": "crawler_registry",
                "required_for_answering": True,
                "source_hash": f"sha256:{file_sha256(CRAWLER_DB_PATH)}",
                "parse_status": "empty_registry",
            }
        ]
        encoded = json.dumps(records, sort_keys=True, separators=(",", ":")).encode()
        return {
            "inventory_hash": f"sha256:{hashlib.sha256(encoded).hexdigest()}",
            "active_source_count": 0,
            "unresolved_count": 1,
            "unresolved_required_count": 1,
            "unresolved_optional_count": 0,
            "unresolved": records,
            "registry_status": "empty",
        }
    records: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    for path in _active_registry_files(
        source_files,
        selected_file=selected_file,
        selected_source_type=selected_source_type,
    ):
        source_type = source_type_for_file(path)
        family = SOURCE_FAMILY_MAP.get(source_type)
        failure = _inspect_parse_artifacts(path)
        json_path, markdown_path, metadata_path = parsed_output_paths(path, source_type)
        record = {
            "source_file": normalize_local_path(str(path)),
            "source_type": source_type,
            "required_for_answering": (
                family.required_for_answering if family is not None else True
            ),
            "source_hash": _existing_file_hash(path),
            "parse_status": failure.get("parse_status") if failure else "parsed",
            "artifact_hashes": {
                "json": _existing_file_hash(json_path),
                "markdown": _existing_file_hash(markdown_path),
                "metadata": _existing_file_hash(metadata_path),
            },
        }
        records.append(record)
        if failure:
            unresolved.append(record)
    encoded = json.dumps(records, sort_keys=True, separators=(",", ":")).encode()
    return {
        "inventory_hash": f"sha256:{hashlib.sha256(encoded).hexdigest()}",
        "active_source_count": len(records),
        "unresolved_count": len(unresolved),
        "unresolved_required_count": sum(
            1 for record in unresolved if record["required_for_answering"]
        ),
        "unresolved_optional_count": sum(
            1 for record in unresolved if not record["required_for_answering"]
        ),
        "unresolved": unresolved,
    }


def _write_run_state(state: dict) -> Path:
    state_path = PIPELINE_DATA_DIR / "run-state" / "latest-parse.json"
    state_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = state_path.with_suffix(".tmp")
    temporary.write_text(json.dumps(state, indent=2, default=str))
    os.replace(temporary, state_path)
    return state_path


def parse_all(
    source_type: str | None = None,
    force: bool = False,
    file_path: Path | None = None,
    *,
    retry_failures: bool = False,
    limit: int | None = None,
    max_runtime_seconds: int | None = None,
    max_consecutive_timeouts: int | None = None,
) -> dict:
    """Parse selected source files and persist an atomic, resumable run summary."""
    converter: DocumentConverter | None = None
    parse_timeout_seconds = get_parse_timeout_seconds()
    started = time.monotonic()
    started_at = datetime.now(timezone.utc).isoformat()
    terminal_state = "running"
    stats = {
        "parsed": 0,
        "skipped": 0,
        "errors": 0,
        "timed_out": 0,
        "quarantined": 0,
        "remaining_failures": 0,
        "remaining_required_failures": 0,
        "remaining_optional_failures": 0,
        "terminal_state": terminal_state,
    }

    if file_path:
        if file_path.suffix.lower() in PARSE_SUPPORTED_FILE_FORMATS:
            all_source_files = [file_path]
            unsupported_files: list[Path] = []
        else:
            all_source_files = []
            unsupported_files = [file_path]
    else:
        all_source_files = discover_source_files(source_type)
        unsupported_files = discover_unsupported_files(source_type)

    source_files = all_source_files
    selected_inventory_file = (
        file_path
        if file_path and file_path.suffix.lower() in PARSE_SUPPORTED_FILE_FORMATS
        else None
    )
    if retry_failures:
        failures_by_path = {
            Path(failure["source_file"]): failure
            for failure in inventory_parse_failures(
                source_files,
                selected_file=selected_inventory_file,
                selected_source_type=source_type,
            )
            if failure.get("retryable", True)
        }
        source_files = sorted(
            failures_by_path,
            key=lambda path: (
                failures_by_path[path].get("attempt_count_generation", 0),
                str(path),
            ),
        )
    if limit is not None:
        source_files = source_files[: max(0, limit)]

    run_state = {
        "started_at": started_at,
        "status": "running",
        "source_type": source_type,
        "retry_failures": retry_failures,
        "selected_file_count": len(source_files),
        "total_file_count": len(all_source_files),
        "processed_file_count": 0,
        "last_source_file": None,
    }
    _write_run_state(run_state)

    try:
        for unsupported_path in unsupported_files:
            source_type_name = source_type_for_file(unsupported_path)
            save_unsupported_format(unsupported_path, source_type_name)
            stats["quarantined"] += 1
            console.print(
                f"  Quarantined {unsupported_path.name} "
                f"(unsupported format {unsupported_path.suffix.lower() or '(none)'})"
            )

        console.print(
            f"[bold]Found {len(source_files)} supported source files to process[/bold]"
        )
        consecutive_timeouts = 0
        total_files = len(source_files)
        for index, source_path in enumerate(source_files, start=1):
            worker_timeout = parse_timeout_seconds
            if max_runtime_seconds is not None:
                remaining_runtime = max_runtime_seconds - (time.monotonic() - started)
                available_worker_time = (
                    remaining_runtime - WORKER_CLEANUP_RESERVE_SECONDS
                )
                if available_worker_time <= 0:
                    terminal_state = "max_runtime_reached"
                    break
                worker_timeout = min(
                    parse_timeout_seconds, max(1, int(available_worker_time))
                )

            source_type_name = source_type_for_file(source_path)
            _, _, meta_path = parsed_output_paths(source_path, source_type_name)
            run_state["last_source_file"] = str(source_path)

            if not force and not should_parse(source_path, meta_path):
                stats["skipped"] += 1
                run_state["processed_file_count"] += 1
                run_state["stats"] = {**stats, "terminal_state": "running"}
                _write_run_state(run_state)
                console.print(
                    f"  [{index}/{total_files}] Skipping {source_path.name} (cached)"
                )
                continue

            try:
                console.print(
                    f"  [{index}/{total_files}] Parsing {source_path.name}...", end=" "
                )
                if source_path.suffix.lower() == ".pdf" and parse_timeout_seconds > 0:
                    parsed_summary = parse_and_save_supported_file_with_timeout(
                        source_path, source_type_name, worker_timeout
                    )
                else:
                    converter_to_use: DocumentConverter | None = None
                    if source_path.suffix.lower() == ".pdf":
                        if converter is None:
                            converter = create_converter()
                        converter_to_use = converter
                    parsed_summary = parse_and_save_supported_file(
                        source_path, source_type_name, converter_to_use
                    )
                console.print(
                    f"[green]OK[/green] ({parsed_summary['section_count']} sections, "
                    f"{parsed_summary['table_count']} tables, "
                    f"{parsed_summary['parse_time_seconds']}s)"
                )
                stats["parsed"] += 1
                consecutive_timeouts = 0
            except Exception as error:
                _, _, committed_meta = parsed_output_paths(
                    source_path, source_type_name
                )
                if committed_meta.exists() and not should_parse(
                    source_path, committed_meta
                ):
                    stats["parsed"] += 1
                    consecutive_timeouts = 0
                    console.print(
                        "[green]OK[/green] "
                        "(artifacts committed before worker IPC closed)"
                    )
                    continue
                save_parse_error(source_path, source_type_name, error)
                console.print(f"[red]ERROR: {error}[/red]")
                stats["errors"] += 1
                if isinstance(error, TimeoutError):
                    stats["timed_out"] += 1
                    consecutive_timeouts += 1
                else:
                    consecutive_timeouts = 0
            finally:
                run_state["processed_file_count"] += 1
                run_state["stats"] = {**stats, "terminal_state": "running"}
                _write_run_state(run_state)

            if (
                max_consecutive_timeouts is not None
                and max_consecutive_timeouts > 0
                and consecutive_timeouts >= max_consecutive_timeouts
            ):
                terminal_state = "max_consecutive_timeouts_reached"
                break
        else:
            terminal_state = "completed"
    except BaseException:
        terminal_state = "interrupted"
        raise
    finally:
        failures = inventory_parse_failures(
            all_source_files,
            selected_file=selected_inventory_file,
            selected_source_type=source_type,
        )
        stats["remaining_failures"] = len(failures)
        required_failures = [
            failure for failure in failures if failure["required_for_answering"]
        ]
        optional_failures = [
            failure for failure in failures if not failure["required_for_answering"]
        ]
        stats["remaining_required_failures"] = len(required_failures)
        stats["remaining_optional_failures"] = len(optional_failures)
        stats["inventory_evidence"] = build_parse_inventory_evidence(
            all_source_files,
            selected_file=selected_inventory_file,
            selected_source_type=source_type,
        )
        stats["terminal_state"] = terminal_state
        run_state.update(
            {
                "status": terminal_state,
                "finished_at": datetime.now(timezone.utc).isoformat(),
                "stats": stats,
                "remaining_failures": failures,
                "remaining_required_failures": required_failures,
                "remaining_optional_failures": optional_failures,
            }
        )
        _write_run_state(run_state)

    return stats
