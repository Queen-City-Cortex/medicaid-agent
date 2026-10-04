"""Native-coordinate proofs for trailing empty fee-schedule XLSX cells."""
from __future__ import annotations

import hashlib
import json
import posixpath
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter

FEE_XLSX_REPRESENTATION = "fee-schedule-trailing-empty-v1"
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def representation_for(source: Path, family: str) -> str | None:
    return FEE_XLSX_REPRESENTATION if family == "fee_schedule" and source.suffix.lower() == ".xlsx" else None


def native_layout(source: Path) -> dict:
    """Read formula/cache coordinates and merges without calculating formulas."""
    layouts = {}
    with ZipFile(source) as archive:
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook.find(f"{MAIN}sheets"):
            target = targets[sheet.attrib[f"{REL}id"]]
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            formulas, merges = [], []
            last_cached_value_row = 0
            row_index, column_index = 0, 0
            with archive.open(path) as stream:
                for event, element in ET.iterparse(stream, events=("start", "end")):
                    if event == "start":
                        if element.tag == f"{MAIN}row":
                            row_index = int(element.attrib.get("r", row_index + 1))
                            column_index = 0
                        continue
                    if element.tag == f"{MAIN}c":
                        cell = element.attrib.get("r")
                        column_index = coordinate_to_tuple(cell)[1] if cell else column_index + 1
                        raw_cell = cell
                        cell = f"{get_column_letter(column_index)}{row_index}"
                        cached_value = element.find(f"{MAIN}v")
                        inline_value = element.find(f"{MAIN}is")
                        if (cached_value is not None and cached_value.text is not None) or inline_value is not None:
                            last_cached_value_row = max(last_cached_value_row, row_index)
                        if element.find(f"{MAIN}f") is not None:
                            cached = element.find(f"{MAIN}v")
                            formulas.append({"cell": cell, "missing_cache": cached is None or cached.text is None, **({"xml_cell": raw_cell} if raw_cell and raw_cell != cell else {})})
                        element.clear()
                    elif element.tag == f"{MAIN}mergeCell":
                        merges.append(element.attrib["ref"])
                        element.clear()
                    elif element.tag == f"{MAIN}row":
                        element.clear()
            layouts[sheet.attrib["name"]] = {"formulas": formulas, "merged_ranges": merges, "last_cached_value_row": last_cached_value_row}
    return layouts


def render_sheet(sheet, layout: dict) -> tuple[str, dict]:
    """Emit original row prefixes and prove populated records stay at their coordinates."""
    protected = {}
    for formula in layout["formulas"]:
        row, column = coordinate_to_tuple(formula["cell"])
        protected[row] = max(protected.get(row, 0), column)
    merges = [range_boundaries(value) for value in layout["merged_ranges"]]
    original_hash, original_cells, emitted_cells = (hashlib.sha256() for _ in range(3))
    rows, widths = [], []
    original_bytes = 0
    native_last_row = layout["last_cached_value_row"]
    original_max_row, original_max_column = sheet.max_row, sheet.max_column
    rendered_last_row = min(original_max_row or native_last_row, native_last_row)
    rows_to_read = sheet.iter_rows(values_only=True, max_row=rendered_last_row) if rendered_last_row else []
    for row_index, row in enumerate(rows_to_read, 1):
        values = [str(value).strip() if value is not None else "" for value in row]
        if not any(values):
            widths.append([row_index, len(values), 0])
            continue
        protected_width = protected.get(row_index, 0)
        for _left, top, right, bottom in merges:
            if top <= row_index <= bottom:
                protected_width = max(protected_width, right)
        if protected_width > len(values):
            raise ValueError("Protected cell lies outside the original rendered row")
        width = len(values)
        while width > protected_width and not values[width - 1]:
            width -= 1
        emitted = values[:width]
        baseline = " | ".join(values).encode("utf-8")
        if rows:
            original_hash.update(b"\n")
            original_bytes += 1
        original_hash.update(baseline)
        original_bytes += len(baseline)
        for column, value in enumerate(values, 1):
            if value:
                original_cells.update((json.dumps([row_index, column, value], ensure_ascii=False) + "\n").encode())
        for column, value in enumerate(emitted, 1):
            if value:
                emitted_cells.update((json.dumps([row_index, column, value], ensure_ascii=False) + "\n").encode())
        if any(values[width:]) or emitted != values[:width]:
            raise ValueError("Candidate row is not an empty-tail-only prefix")
        rows.append(" | ".join(emitted))
        widths.append([row_index, len(values), width])
    # Independently read all native rows, ignoring potentially understated dimensions.
    # Cached values retain the existing openpyxl conversion, including dates/booleans.
    native_cells = hashlib.sha256()
    sheet.reset_dimensions()
    native_rows = sheet.iter_rows(values_only=True, max_row=native_last_row) if native_last_row else []
    for row_index, row in enumerate(native_rows, 1):
        for column, value in enumerate(row, 1):
            normalized = str(value).strip() if value is not None else ""
            if normalized:
                native_cells.update((json.dumps([row_index, column, normalized], ensure_ascii=False) + "\n").encode())
    if native_cells.digest() != original_cells.digest():
        raise ValueError("Worksheet dimensions omit populated native cells")
    if original_cells.digest() != emitted_cells.digest():
        raise ValueError("Populated native cell records changed")
    body = "\n".join(rows)
    return body, {
        **layout, "rows": widths,
        "omitted_empty_tail_rows": [rendered_last_row + 1, original_max_row, original_max_column] if original_max_row and original_max_row > rendered_last_row else None, "baseline_body_sha256": original_hash.hexdigest(),
        "baseline_utf8_bytes": original_bytes, "candidate_utf8_bytes": len(body.encode()),
        "candidate_body_sha256": hashlib.sha256(body.encode()).hexdigest(),
        "native_populated_cells_sha256": native_cells.hexdigest(),
        "emitted_populated_cells_sha256": emitted_cells.hexdigest(),
        "prefix_and_empty_tail_verified": True,
    }
