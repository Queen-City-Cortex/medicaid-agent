"""Native cell and targeted cache regressions (large fixtures have short IDs)."""
import hashlib
import json
from datetime import datetime
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from pipeline.parse import converter
from pipeline.parse.fee_xlsx import FEE_XLSX_REPRESENTATION


def workbook(tmp_path):
    root = tmp_path / "fee-schedules"
    root.mkdir()
    path = root / "fees.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Header"
    sheet["B1"] = "Merged"
    sheet.merge_cells("B1:F1")
    sheet["B3"] = 0
    sheet["D3"] = False
    sheet["B4"] = "λ | value\tand\nnewline"
    sheet["D4"] = 0.123456789012345
    sheet["F4"] = datetime(2021, 1, 2)
    sheet["A5"] = "formula row"
    sheet["H5"] = "=1+2"
    sheet["H6"] = "=1+2"  # Cached entirely empty row stays omitted.
    sheet["XFC1"].fill = PatternFill("solid", fgColor="FF0000")
    book.save(path)
    book.close()
    return path


def test_native_prefix_proof_preserves_alignment_and_values(tmp_path):
    path = workbook(tmp_path)
    baseline = converter.parse_xlsx(path)
    candidate = converter.parse_xlsx(path, "fee_schedule")
    payload = candidate["payload"]
    proof = payload["cell_preservation"][0]
    original = baseline["payload"]["synthetic_sections"][0]["content_text"]
    body = payload["synthetic_sections"][0]["content_text"]
    assert proof["baseline_body_sha256"] == hashlib.sha256(original.encode()).hexdigest()
    assert proof["native_populated_cells_sha256"] == proof["emitted_populated_cells_sha256"]
    assert proof["rows"][:3] == [[1, 16383, 6], [2, 16383, 0], [3, 16383, 4]]
    assert proof["rows"][4:] == [[5, 16383, 8], [6, 16383, 0]]
    assert proof["formulas"] == [{"cell": "H5", "missing_cache": True}, {"cell": "H6", "missing_cache": True}]
    assert body.startswith(" | Merged |  |  |  | \n | 0 |  | False\n")
    assert "λ | value\tand\nnewline" in body
    assert "0.123456789012345" in body
    assert "2021-01-02 00:00:00" in body
    assert body.endswith("formula row" + " | " * 7)
    assert len(body) < len(original) / 100
    assert baseline["payload"].get("representation") is None


def test_formula_with_cached_value_is_not_recalculated(tmp_path):
    path = workbook(tmp_path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>999</v>", 1)
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    parsed = converter.parse_supported_file(path, None, "fee_schedule")
    assert parsed["payload"]["cell_preservation"][0]["formulas"][0]["missing_cache"] is False
    assert parsed["payload"]["synthetic_sections"][0]["content_text"].endswith("999")


@pytest.mark.parametrize("version", [4, 5, 6], ids=["v4", "v5", "v6"])
def test_old_fee_cache_never_migrated_or_counted_as_same_generation(monkeypatch, tmp_path, version):
    path = workbook(tmp_path)
    monkeypatch.setattr(converter, "PARSED_DATA_DIR", tmp_path / "parsed")
    parsed = converter.parse_xlsx(path, "fee_schedule")
    converter.save_parsed(path, parsed, "fee_schedule")
    _, _, meta_path = converter.parsed_output_paths(path, "fee_schedule")
    meta = json.loads(meta_path.read_text())
    meta.pop("representation")
    meta["parse_cache_version"] = version
    meta["attempt_count_generation"] = 9
    meta["attempt_count_total"] = 11
    meta_path.write_text(json.dumps(meta))
    before = meta_path.read_bytes()
    assert converter.should_parse(path, meta_path)
    assert meta_path.read_bytes() == before
    assert converter._inspect_parse_artifacts(path)["parse_status"] == "stale_sidecar"
    converter.save_parsed(path, parsed, "fee_schedule")
    meta = json.loads(meta_path.read_text())
    assert (meta["attempt_count_generation"], meta["attempt_count_total"]) == (1, 12)
    assert meta["representation"] == FEE_XLSX_REPRESENTATION
    assert not converter.should_parse(path, meta_path)
    assert converter._inspect_parse_artifacts(path) is None


def test_save_rejects_old_representation_and_worker_passes_family(monkeypatch, tmp_path):
    path = workbook(tmp_path)
    monkeypatch.setattr(converter, "PARSED_DATA_DIR", tmp_path / "parsed")
    with pytest.raises(ValueError, match="representation"):
        converter.save_parsed(path, converter.parse_xlsx(path), "fee_schedule")
    results = []
    class Connection:
        def send(self, result):
            results.append(result)
    converter._parse_supported_file_worker(str(path), "fee_schedule", Connection())
    assert results[0]["ok"]
    _, _, meta = converter.parsed_output_paths(path, "fee_schedule")
    assert not converter.should_parse(path, meta)


def test_understated_dimension_rejected_against_native_cells(tmp_path):
    path = workbook(tmp_path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'ref="A1:XFC6"', b'ref="A1:D1"')
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="Protected cell|dimensions omit"):
        converter.parse_xlsx(path, "fee_schedule")


def test_dimension_omits_populated_unprotected_values(tmp_path):
    path = tmp_path / "understated.xlsx"
    book = Workbook()
    book.active["A1"] = "visible"
    book.active["B1"] = "hidden column"
    book.active["A2"] = "hidden row"
    book.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'ref="A1:B2"', b'ref="A1:A1"')
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="dimensions omit"):
        converter.parse_xlsx(path, "fee_schedule")


def test_error_recovery_resets_attempts_for_targeted_representation(monkeypatch, tmp_path):
    path = workbook(tmp_path)
    monkeypatch.setattr(converter, "PARSED_DATA_DIR", tmp_path / "parsed")
    converter.save_parse_error(path, "fee_schedule", ValueError("old error"))
    _, _, meta_path = converter.parsed_output_paths(path, "fee_schedule")
    meta = json.loads(meta_path.read_text())
    meta.pop("representation")
    meta["attempt_count_generation"] = 15
    meta_path.write_text(json.dumps(meta))
    assert converter._inspect_parse_artifacts(path)["attempt_count_generation"] == 0
    converter.save_parse_error(path, "fee_schedule", ValueError("new error"))
    meta = json.loads(meta_path.read_text())
    assert meta["attempt_count_generation"] == 1
    assert converter._inspect_parse_artifacts(path)["parse_status"] == "parse_exception"


def test_vertical_merges_offset_cells_and_nonfee_unchanged(tmp_path):
    source = tmp_path / "offset.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["C3"] = "Offset header"
    sheet.merge_cells("C3:G5")
    sheet["A4"] = "Second row"
    sheet["K6"].number_format = "0.00"
    book.save(source)
    book.close()
    baseline = converter.parse_xlsx(source)
    untouched = converter.parse_supported_file(source, None, "clinical_policy")
    compact = converter.parse_supported_file(source, None, "fee_schedule")
    assert untouched["payload"] == baseline["payload"]
    proof = compact["payload"]["cell_preservation"][0]
    assert proof["rows"] == [[1, 11, 0], [2, 11, 0], [3, 11, 7], [4, 11, 7], [5, 11, 0], [6, 11, 0]]
    assert compact["payload"]["synthetic_sections"][0]["content_text"].startswith(" |  | Offset header")


def test_implicit_formula_coordinate_and_shared_formula_followers(tmp_path):
    source = workbook(tmp_path)
    with ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    data = entries["xl/worksheets/sheet1.xml"]
    # Make H5 implicit: preceding G5 establishes its column without an r attribute.
    data = data.replace(b'<c r="H5"><f>1+2</f>', b'<c r="G5" t="n"><v></v></c><c><f t="shared" si="0" ref="H5:H6">1+2</f>')
    data = data.replace(b'<c r="H6"><f>1+2</f>', b'<c r="H6"><f t="shared" si="0"/>')
    entries["xl/worksheets/sheet1.xml"] = data
    with ZipFile(source, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    parsed = converter.parse_xlsx(source, "fee_schedule")
    assert [item["cell"] for item in parsed["payload"]["cell_preservation"][0]["formulas"]] == ["H5", "H6"]


def test_merged_span_outside_original_dimensions_fails_explicitly(tmp_path):
    source = tmp_path / "merge.xlsx"
    book = Workbook()
    book.active["A1"] = "Header"
    book.active.merge_cells("A1:D1")
    book.save(source)
    with ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'<dimension ref="A1:D1"/>', b'')
    with ZipFile(source, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="Protected cell lies outside"):
        converter.parse_xlsx(source, "fee_schedule")
