"""Tests for supported source-file parsing beyond PDFs."""

import json
import subprocess
import sys
import textwrap
from pathlib import Path

import pytest
from docling_core.types.doc import DocItemLabel
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from pipeline.parse import converter as parse_converter
from pipeline.parse import worker as parse_worker
from pipeline.parse.converter import (
    PARSE_CACHE_VERSION,
    artifact_stem_for_source_path,
    build_parse_inventory_evidence,
    create_converter,
    evaluate_no_ocr_pdf,
    inventory_parse_failures,
    parse_all,
    parse_and_save_supported_file_with_timeout,
    parse_doc,
    parse_docx,
    parse_html,
    parse_pdf,
    parse_pptx,
    parse_xls,
    parse_xlsx,
    save_parse_error,
    save_parsed,
    should_parse,
    _validate_worker_result,
)


class _PreflightPage(dict):
    def __init__(self, text, *, image=False):
        super().__init__()
        self.text = text
        self.image = image

    def extract_text(self):
        return self.text


def _mock_preflight(monkeypatch, pages):
    class Reader:
        def __init__(self, *_args, **_kwargs):
            self.pages = pages

    monkeypatch.setattr("pipeline.parse.converter.PdfReader", Reader)
    monkeypatch.setattr(
        "pipeline.parse.converter._page_has_images", lambda page: page.image
    )


def test_no_ocr_preflight_rejects_mixed_digital_and_scanned(monkeypatch, tmp_path):
    digital = "eligibility requirements and covered services " * 4
    _mock_preflight(
        monkeypatch, [_PreflightPage(digital), _PreflightPage("", image=True)]
    )
    accepted, reason, evidence = evaluate_no_ocr_pdf(
        tmp_path / "mixed.pdf", {"page_content": {1: digital}}
    )
    assert accepted is False
    assert reason == "page_2_image_or_sparse_source"
    assert evidence[1]["has_images"] is True


def test_no_ocr_preflight_rejects_repeated_header_only(monkeypatch, tmp_path):
    _mock_preflight(monkeypatch, [_PreflightPage("NC Medicaid Header")])
    accepted, reason, _ = evaluate_no_ocr_pdf(
        tmp_path / "header.pdf", {"page_content": {1: "NC Medicaid Header"}}
    )
    assert accepted is False
    assert reason == "page_1_image_or_sparse_source"


def test_no_ocr_preflight_accepts_populated_page_then_true_blank(monkeypatch, tmp_path):
    text = "provider enrollment obligations and appeal procedures " * 4
    _mock_preflight(monkeypatch, [_PreflightPage(text), _PreflightPage("")])
    accepted, reason, evidence = evaluate_no_ocr_pdf(
        tmp_path / "blank-tail.pdf", {"page_content": {1: text}}
    )
    assert accepted is True
    assert reason.startswith("all_nonblank_pages")
    assert evidence[1]["source_blank"] is True


def test_no_ocr_preflight_requires_ocr_for_substantive_tables(monkeypatch, tmp_path):
    table = "procedure code description rate effective date " * 4
    _mock_preflight(monkeypatch, [_PreflightPage(table)])
    accepted, reason, evidence = evaluate_no_ocr_pdf(
        tmp_path / "table.pdf", {"page_content": {1: table}, "table_count": 1}
    )
    assert accepted is False
    assert reason == "table_structure_requires_ocr"
    assert evidence[0]["normalized_token_coverage"] == 1.0


def test_no_ocr_preflight_rejects_corrupt_or_sparse_text_layer(monkeypatch, tmp_path):
    source = "covered benefit prior authorization clinical criteria " * 4
    _mock_preflight(monkeypatch, [_PreflightPage(source)])
    accepted, reason, _ = evaluate_no_ocr_pdf(
        tmp_path / "corrupt.pdf",
        {"page_content": {1: "completely unrelated extracted material " * 4}},
    )
    assert accepted is False
    assert reason == "page_1_insufficient_text_coverage"


def test_no_ocr_preflight_rejects_missing_repeated_material(monkeypatch, tmp_path):
    repeated = "covered services require prior authorization "
    source = repeated * 12
    _mock_preflight(monkeypatch, [_PreflightPage(source)])

    accepted, reason, evidence = evaluate_no_ocr_pdf(
        tmp_path / "repeated.pdf",
        {"page_content": {1: repeated}},
    )

    assert accepted is False
    assert reason == "page_1_insufficient_text_coverage"
    assert evidence[0]["normalized_token_coverage"] < 0.1


def test_no_ocr_preflight_rejects_reordered_text(monkeypatch, tmp_path):
    source = "coverage is denied unless prior authorization is approved " * 4
    reordered = " ".join(reversed(source.split()))
    _mock_preflight(monkeypatch, [_PreflightPage(source)])

    accepted, reason, evidence = evaluate_no_ocr_pdf(
        tmp_path / "reordered.pdf",
        {"page_content": {1: reordered}, "table_count": 0},
    )

    assert accepted is False
    assert reason == "page_1_insufficient_text_coverage"
    assert evidence[0]["normalized_token_sequence_matches"] is False


def test_worker_result_schema_rejects_malformed_and_caps_error():
    with pytest.raises(RuntimeError, match="malformed"):
        _validate_worker_result({"ok": "yes", "summary": {}})
    with pytest.raises(RuntimeError, match="malformed success"):
        _validate_worker_result({"ok": True, "summary": {"section_count": 1}})
    result = _validate_worker_result({"ok": False, "error_message": "x" * 10000})
    assert len(result["error_message"]) == 4096


def _parsed_fixture():
    return {
        "payload": {"name": "fixture"},
        "markdown": "substantive fixture markdown",
        "source_format": "xlsx",
        "parse_backend": "test",
        "parse_status": "parsed",
        "quality_flags": [],
        "answer_ready": True,
        "section_count": 1,
        "table_count": 0,
        "page_count": 1,
        "content_char_count": 30,
        "parse_time_seconds": 0.1,
    }


def test_atomic_artifact_hashes_force_retry_after_truncation_or_deletion(
    monkeypatch, tmp_path
):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"generation-one")
    paths = (
        tmp_path / "source.json",
        tmp_path / "source.md",
        tmp_path / "source.meta.json",
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths", lambda *_args: paths
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file", lambda _path: "fee_schedule"
    )

    save_parsed(source, _parsed_fixture(), "fee_schedule")
    meta = json.loads(paths[2].read_text())
    assert meta["artifacts"]["json"]["sha256"].startswith("sha256:")
    assert should_parse(source, paths[2]) is False

    paths[0].write_text("{")
    assert should_parse(source, paths[2]) is True
    save_parsed(source, _parsed_fixture(), "fee_schedule")
    paths[1].unlink()
    assert should_parse(source, paths[2]) is True


def test_attempt_counts_reset_by_generation_but_preserve_total(monkeypatch, tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"first")
    meta_path = tmp_path / "source.meta.json"
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda *_args: (tmp_path / "source.json", tmp_path / "source.md", meta_path),
    )
    save_parse_error(source, "clinical_policy", RuntimeError("one"))
    save_parse_error(source, "clinical_policy", RuntimeError("two"))
    assert json.loads(meta_path.read_text())["attempt_count_generation"] == 2

    source.write_bytes(b"second")
    save_parse_error(source, "clinical_policy", RuntimeError("new generation"))
    meta = json.loads(meta_path.read_text())
    assert meta["attempt_count_generation"] == 1
    assert meta["attempt_count_total"] == 3


class _FakePipeEnd:
    def __init__(self, events, *, result=None, readable=False, eof=False):
        self.events = events
        self.result = result
        self.readable = readable
        self.eof = eof

    def poll(self, _timeout=0):
        self.events.append("poll")
        return self.readable

    def recv(self):
        self.events.append("recv")
        if self.eof:
            raise EOFError
        return self.result

    def close(self):
        self.events.append("close")


class _FakeProcess:
    def __init__(self, events, *, alive):
        self.events = events
        self.alive = alive
        self.exitcode = 7 if not alive else None

    def start(self):
        self.events.append("start")

    def is_alive(self):
        return self.alive

    def terminate(self):
        self.events.append("terminate")
        self.alive = False
        self.exitcode = -15

    def kill(self):
        self.events.append("kill")
        self.alive = False
        self.exitcode = -9

    def join(self, *_args):
        self.events.append("join")


def _install_fake_context(
    monkeypatch, *, result=None, readable=False, alive=True, eof=False
):
    events = []
    parent = _FakePipeEnd(events, result=result, readable=readable, eof=eof)
    child = _FakePipeEnd(events)
    process = _FakeProcess(events, alive=alive)

    class Context:
        def Pipe(self, **_kwargs):
            return parent, child

        def Process(self, **_kwargs):
            return process

    monkeypatch.setattr(
        "pipeline.parse.converter.get_context", lambda *_args: Context()
    )
    return events


def test_worker_parent_receives_before_join_for_pipe_payload(monkeypatch, tmp_path):
    summary = {"section_count": 1, "table_count": 2, "parse_time_seconds": 0.1}
    events = _install_fake_context(
        monkeypatch, result={"ok": True, "summary": summary}, readable=True
    )
    assert (
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 1)
        == summary
    )
    assert events.index("recv") < events.index("join")


def test_worker_success_gets_natural_exit_grace_without_forced_cleanup(
    monkeypatch, tmp_path
):
    events = []
    summary = {"section_count": 1, "table_count": 0, "parse_time_seconds": 0.1}
    parent = _FakePipeEnd(
        events,
        result={"ok": True, "summary": summary},
        readable=True,
    )
    child = _FakePipeEnd(events)

    class NaturalExitProcess(_FakeProcess):
        def join(self, *_args):
            self.events.append("join")
            self.alive = False
            self.exitcode = 0

    process = NaturalExitProcess(events, alive=True)

    class Context:
        def Pipe(self, **_kwargs):
            return parent, child

        def Process(self, **_kwargs):
            return process

    monkeypatch.setattr(parse_converter, "get_context", lambda *_args: Context())
    cleaned = []
    monkeypatch.setattr(
        parse_converter,
        "_cleanup_killed_worker_semaphores",
        cleaned.append,
    )

    assert (
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 1)
        == summary
    )
    assert "terminate" not in events
    assert "kill" not in events
    assert cleaned == []


@pytest.mark.parametrize(
    ("result", "readable", "alive", "eof", "message"),
    [
        (None, False, False, False, "without result"),
        (None, True, False, True, "without result"),
        ({"unexpected": True}, True, False, False, "malformed"),
    ],
)
def test_worker_handles_early_exit_eof_and_malformed_payload(
    monkeypatch, tmp_path, result, readable, alive, eof, message
):
    _install_fake_context(
        monkeypatch, result=result, readable=readable, alive=alive, eof=eof
    )
    with pytest.raises(RuntimeError, match=message):
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 1)


def test_worker_timeout_repeatedly_closes_and_reaps(monkeypatch, tmp_path):
    for _ in range(5):
        events = _install_fake_context(monkeypatch, readable=False, alive=True)
        with pytest.raises(TimeoutError):
            parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 0)
        assert "terminate" in events
        assert "join" in events
        assert events.count("close") == 2


def test_worker_timeout_unlinks_spawn_semaphores_created_by_killed_child(
    monkeypatch, tmp_path
):
    cleaned_prefixes = []
    monkeypatch.setattr(
        parse_converter,
        "_cleanup_killed_worker_semaphores",
        cleaned_prefixes.append,
    )
    _install_fake_context(monkeypatch, readable=False, alive=True)

    with pytest.raises(TimeoutError):
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 0)

    assert len(cleaned_prefixes) == 1
    assert cleaned_prefixes[0].startswith("/pipeline-pdf-")


def test_worker_cleanup_does_not_unlink_namespace_until_worker_is_reaped(
    monkeypatch, tmp_path
):
    events = []
    parent = _FakePipeEnd(events)
    child = _FakePipeEnd(events)

    class UnreapableProcess(_FakeProcess):
        def terminate(self):
            self.events.append("terminate")

        def kill(self):
            self.events.append("kill")

    process = UnreapableProcess(events, alive=True)

    class Context:
        def Pipe(self, **_kwargs):
            return parent, child

        def Process(self, **_kwargs):
            return process

    monkeypatch.setattr(parse_converter, "get_context", lambda *_args: Context())
    cleaned = []
    monkeypatch.setattr(
        parse_converter,
        "_cleanup_killed_worker_semaphores",
        cleaned.append,
    )

    with pytest.raises(TimeoutError):
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 0)

    assert "terminate" in events
    assert "kill" in events
    assert cleaned == []


@pytest.mark.skipif(sys.platform != "linux", reason="requires Linux /dev/shm semantics")
def test_killed_spawn_worker_exits_without_resource_tracker_semaphore_warning(tmp_path):
    script = tmp_path / "semaphore_cleanup_repro.py"
    ready = tmp_path / "ready"
    script.write_text(
        textwrap.dedent(
            """
            import multiprocessing as mp
            import sys
            import time
            from pathlib import Path

            from pipeline.parse.converter import (
                _cleanup_killed_worker_semaphores,
            )
            from pipeline.parse.worker import _configure_semaphore_prefix


            def worker(prefix, ready_path):
                _configure_semaphore_prefix(prefix)
                semaphore = mp.Semaphore(1)
                Path(ready_path).write_text(semaphore._semlock.name)
                time.sleep(30)


            if __name__ == "__main__":
                prefix_a = "/pipeline-test-a"
                prefix_b = "/pipeline-test-b"
                ready_a = Path(sys.argv[1])
                ready_b = ready_a.with_name("ready-b")
                context = mp.get_context("spawn")
                default_semaphore = context.Semaphore(1)
                default_path = Path(
                    "/dev/shm/sem." + default_semaphore._semlock.name.lstrip("/")
                )
                process_a = context.Process(
                    target=worker,
                    args=(prefix_a, ready_a),
                )
                process_b = context.Process(
                    target=worker,
                    args=(prefix_b, ready_b),
                )
                process_a.start()
                process_b.start()
                deadline = time.monotonic() + 15
                while not (ready_a.exists() and ready_b.exists()):
                    if time.monotonic() >= deadline:
                        raise RuntimeError("workers did not create semaphores")
                    time.sleep(0.05)
                semaphore_b_path = Path(
                    "/dev/shm/sem." + ready_b.read_text().lstrip("/")
                )
                process_a.terminate()
                process_a.join(5)
                _cleanup_killed_worker_semaphores(prefix_a)
                assert semaphore_b_path.exists()
                assert default_path.exists()
                process_b.terminate()
                process_b.join(5)
                _cleanup_killed_worker_semaphores(prefix_b)
            """
        )
    )

    result = subprocess.run(
        [sys.executable, str(script), str(ready)],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )

    assert result.returncode == 0, result.stderr
    assert "leaked semaphore" not in result.stderr


def test_worker_start_failure_preserves_error_and_closes_pipe(monkeypatch, tmp_path):
    events = []
    parent = _FakePipeEnd(events)
    child = _FakePipeEnd(events)

    class StartFailProcess(_FakeProcess):
        def start(self):
            self.events.append("start")
            raise OSError("spawn failed")

    process = StartFailProcess(events, alive=False)

    class Context:
        def Pipe(self, **_kwargs):
            return parent, child

        def Process(self, **_kwargs):
            return process

    monkeypatch.setattr(
        "pipeline.parse.converter.get_context", lambda *_args: Context()
    )

    with pytest.raises(OSError, match="spawn failed"):
        parse_and_save_supported_file_with_timeout(tmp_path / "x.pdf", "test", 1)
    assert "join" not in events
    assert events.count("close") == 2


def test_spawn_worker_reports_bootstrap_errors(monkeypatch):
    queued_results = []

    class RecordingConnection:
        def send(self, result):
            queued_results.append(result)

        def close(self):
            pass

    def fail_bootstrap():
        raise ImportError("bad bootstrap")

    monkeypatch.setattr(parse_worker, "_preload_stdlib_typing", fail_bootstrap)

    parse_worker.parse_supported_file_worker(
        "/tmp/source.xlsx",
        "fee_schedule",
        RecordingConnection(),
    )

    assert queued_results == [
        {
            "ok": False,
            "error_message": "ImportError: bad bootstrap",
        }
    ]


def test_converter_preloads_stdlib_typing_before_docling_imports():
    script = textwrap.dedent(
        """
        import builtins
        import sys

        real_import = builtins.__import__
        sys.modules.pop("typing", None)

        def guarded_import(name, globals=None, locals=None, fromlist=(), level=0):
            if (name == "docling" or name.startswith("docling.")) and "typing" not in sys.modules:
                raise RuntimeError("docling imported before standard-library typing")
            return real_import(name, globals, locals, fromlist, level)

        builtins.__import__ = guarded_import
        import pipeline.parse.converter

        typing_path = getattr(sys.modules["typing"], "__file__", "")
        assert "site-packages/cv2/typing" not in typing_path
        """
    )

    result = subprocess.run(
        [sys.executable, "-c", script],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr


def test_spawn_worker_repairs_an_inherited_opencv_import_path(tmp_path):
    script = textwrap.dedent(
        f"""
        import os
        import sys
        from pathlib import Path

        os.environ["CORPUS_DATA_ROOT"] = {str(tmp_path)!r}

        from openpyxl import Workbook
        from pipeline.parse.converter import parse_and_save_supported_file_with_timeout

        source_path = Path({str(tmp_path / "worker-smoke.xlsx")!r})
        workbook = Workbook()
        workbook.active.append(["Code", "Rate"])
        workbook.active.append(["T1001", "125.00"])
        workbook.save(source_path)

        import cv2

        sys.path.insert(0, str(Path(cv2.__file__).parent))
        result = parse_and_save_supported_file_with_timeout(
            source_path,
            "fee_schedule",
            timeout_seconds=30,
        )
        assert result["section_count"] == 1
        """
    )

    result = subprocess.run(
        [sys.executable, "-c", script],
        check=False,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert result.returncode == 0, result.stderr


def test_artifact_stem_adds_format_suffix_for_non_pdf_sources():
    assert artifact_stem_for_source_path(Path("/tmp/example.pdf")) == "example"
    assert artifact_stem_for_source_path(Path("/tmp/example.docx")) == "example__docx"
    assert artifact_stem_for_source_path(Path("/tmp/example.xlsx")) == "example__xlsx"


def test_parse_docx_extracts_headings_and_paragraph_content(tmp_path):
    source_path = tmp_path / "billing-guide.docx"

    document = WordDocument()
    document.add_heading("1.0 Overview", level=1)
    document.add_paragraph("This document explains the billing workflow.")
    document.save(source_path)

    parsed = parse_docx(source_path)

    assert parsed["source_format"] == "docx"
    assert parsed["section_count"] == 1
    assert parsed["answer_ready"] is True
    assert parsed["payload"]["synthetic_sections"][0]["title"] == "1.0 Overview"
    assert (
        "billing workflow" in parsed["payload"]["synthetic_sections"][0]["content_text"]
    )


def test_parse_xlsx_extracts_sheets_into_synthetic_sections(tmp_path):
    source_path = tmp_path / "fee-schedule.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Code", "Rate"])
    sheet.append(["T1001", "125.00"])
    workbook.save(source_path)

    parsed = parse_xlsx(source_path)

    assert parsed["source_format"] == "xlsx"
    assert parsed["section_count"] == 1
    assert parsed["answer_ready"] is True
    assert parsed["payload"]["synthetic_sections"][0]["title"] == "Rates"
    assert "T1001" in parsed["payload"]["synthetic_sections"][0]["content_text"]


def test_parse_xls_extracts_legacy_workbook_rows(monkeypatch, tmp_path):
    source_path = tmp_path / "legacy-rates.xls"
    source_path.write_bytes(b"legacy workbook")

    class FakeSheet:
        name = "Rates"
        nrows = 2
        ncols = 2

        def cell_value(self, row, column):
            return [["Code", "Rate"], ["T1001", 125.0]][row][column]

    class FakeWorkbook:
        def sheets(self):
            return [FakeSheet()]

        def release_resources(self):
            return None

    monkeypatch.setattr(
        "pipeline.parse.converter.xlrd.open_workbook",
        lambda *_args, **_kwargs: FakeWorkbook(),
    )

    parsed = parse_xls(source_path)

    assert parsed["source_format"] == "xls"
    assert parsed["parse_backend"] == "xlrd"
    assert parsed["answer_ready"] is True
    assert "T1001" in parsed["payload"]["synthetic_sections"][0]["content_text"]


def test_parse_pptx_extracts_slide_text_and_tables(tmp_path):
    source_path = tmp_path / "managed-care.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Provider Requirements"
    text_box = slide.shapes.add_textbox(0, 0, 100, 100)
    text_box.text = "Plans must publish current guidance."
    table = slide.shapes.add_table(2, 2, 0, 0, 100, 100).table
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "PHP"
    table.cell(1, 1).text = "Notice"
    presentation.save(source_path)

    parsed = parse_pptx(source_path)

    assert parsed["source_format"] == "pptx"
    assert parsed["parse_backend"] == "python-pptx"
    assert parsed["table_count"] == 1
    content = parsed["payload"]["synthetic_sections"][0]["content_text"]
    assert "publish current guidance" in content
    assert "PHP | Notice" in content


def test_parse_doc_uses_headless_libreoffice_conversion(monkeypatch, tmp_path):
    source_path = tmp_path / "legacy-manual.doc"
    source_path.write_bytes(b"legacy doc")

    def fake_run(args, **_kwargs):
        output_dir = Path(args[args.index("--outdir") + 1])
        document = WordDocument()
        document.add_heading("Overview", level=1)
        document.add_paragraph("Converted legacy content.")
        document.save(output_dir / "legacy-manual.docx")

        class Completed:
            returncode = 0
            stdout = ""
            stderr = ""

        return Completed()

    monkeypatch.setattr(
        "pipeline.parse.converter.subprocess.run",
        fake_run,
    )

    parsed = parse_doc(source_path)

    assert parsed["source_format"] == "doc"
    assert parsed["parse_backend"] == "libreoffice+python-docx"
    assert "Converted legacy content" in parsed["markdown"]


def test_parse_html_extracts_only_article_content_and_modified_date(tmp_path):
    source_path = tmp_path / "medicaid-is-changing.html"
    source_path.write_text(
        """<!doctype html><html><body><nav>Ignore this menu</nav><article>
        <h1>Medicaid is Changing</h1><p>General beneficiary guidance.</p>
        <h2>Exceptions and exclusions</h2><p>Ask DSS about individual facts.</p>
        <p>This page was last modified on 07/15/2026</p>
        </article><footer>Ignore this footer</footer></body></html>"""
    )

    parsed = parse_html(source_path)

    assert parsed["source_format"] == "html"
    assert parsed["source_modified_date"] == "2026-07-15"
    assert parsed["answer_ready"] is True
    assert [
        section["title"] for section in parsed["payload"]["synthetic_sections"]
    ] == [
        "Medicaid is Changing",
        "Exceptions and exclusions",
    ]
    assert "Ignore this menu" not in parsed["markdown"]
    assert "Ignore this footer" not in parsed["markdown"]


class _FakeDocItem:
    def __init__(self, label, text=""):
        self.label = label
        self.text = text


class _FakeDoc:
    def __init__(self):
        self.pages = [object()]

    def export_to_dict(self):
        return {
            "tables": [
                {
                    "data": {
                        "table_cells": [
                            {
                                "start_row_offset_idx": 0,
                                "start_col_offset_idx": 0,
                                "end_col_offset_idx": 1,
                                "text": "PHP Name",
                            },
                            {
                                "start_row_offset_idx": 0,
                                "start_col_offset_idx": 1,
                                "end_col_offset_idx": 2,
                                "text": "Contract Name",
                            },
                            {
                                "start_row_offset_idx": 1,
                                "start_col_offset_idx": 0,
                                "end_col_offset_idx": 1,
                                "text": "Healthy Blue",
                            },
                            {
                                "start_row_offset_idx": 1,
                                "start_col_offset_idx": 1,
                                "end_col_offset_idx": 2,
                                "text": "Provider Payment Appendix",
                            },
                        ]
                    }
                }
            ]
        }

    def export_to_markdown(self):
        return (
            "| PHP Name | Contract Name |\n| Healthy Blue | Provider Payment Appendix |"
        )

    def iterate_items(self):
        yield _FakeDocItem(DocItemLabel.TABLE), 0


class _FakeConversionResult:
    def __init__(self):
        self.document = _FakeDoc()


class _FakeConverter:
    def convert(self, _path: str):
        return _FakeConversionResult()


def test_parse_pdf_promotes_table_only_content_into_synthetic_sections(
    monkeypatch, tmp_path
):
    source_path = tmp_path / "managed-care-contracts.pdf"
    source_path.write_bytes(b"%PDF-1.4 fake")

    monkeypatch.setattr(
        "pipeline.parse.converter.evaluate_no_ocr_pdf",
        lambda *_args: (True, "complete", [{"page": 1}]),
    )
    parsed = parse_pdf(source_path, _FakeConverter())

    assert parsed["source_format"] == "pdf"
    assert parsed["table_count"] == 1
    assert parsed["section_count"] == 1
    assert parsed["answer_ready"] is True
    assert parsed["quality_flags"] == []
    assert parsed["content_char_count"] > 0
    assert "Healthy Blue" in parsed["markdown"]


def test_should_parse_reprocesses_when_cache_version_changes(tmp_path):
    source_path = tmp_path / "managed-care-contracts.pdf"
    source_path.write_bytes(b"%PDF-1.4 fake")
    meta_path = tmp_path / "managed-care-contracts.meta.json"
    meta_path.write_text(
        json.dumps(
            {
                "source_file_hash": "sha256:not-the-current-hash",
                "parse_cache_version": PARSE_CACHE_VERSION - 1,
            }
        )
    )

    assert should_parse(source_path, meta_path) is True


def test_should_parse_migrates_valid_v4_success_without_reparsing(
    monkeypatch, tmp_path
):
    source_path = tmp_path / "already-good.pdf"
    source_path.write_bytes(b"%PDF-1.4 already parsed")
    json_path = tmp_path / "already-good.json"
    md_path = tmp_path / "already-good.md"
    meta_path = tmp_path / "already-good.meta.json"
    json_path.write_text(json.dumps({"name": "already good"}))
    md_path.write_text("already good parsed content")
    source_hash = (
        f"sha256:{__import__('hashlib').sha256(source_path.read_bytes()).hexdigest()}"
    )
    meta_path.write_text(
        json.dumps(
            {
                "source_file_hash": source_hash,
                "parse_cache_version": 4,
                "parse_status": "parsed",
                "answer_ready": True,
                "attempt_count": 2,
            }
        )
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda *_args: (json_path, md_path, meta_path),
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )

    assert should_parse(source_path, meta_path) is False
    migrated = json.loads(meta_path.read_text())
    assert migrated["parse_cache_version"] == PARSE_CACHE_VERSION
    assert migrated["cache_migrated_from_version"] == 4
    assert migrated["attempt_count_total"] == 2
    assert migrated["artifacts"]["json"]["sha256"].startswith("sha256:")
    assert migrated["artifacts"]["markdown"]["sha256"].startswith("sha256:")
    assert inventory_parse_failures([source_path]) == []


def test_should_parse_retries_non_answer_ready_success(tmp_path):
    source_path = tmp_path / "partial.pdf"
    source_path.write_bytes(b"%PDF-1.4 partial")
    meta_path = tmp_path / "partial.meta.json"
    meta_path.write_text(
        json.dumps(
            {
                "source_file_hash": f"sha256:{__import__('hashlib').sha256(source_path.read_bytes()).hexdigest()}",
                "parse_cache_version": PARSE_CACHE_VERSION,
                "parse_status": "parsed",
                "answer_ready": False,
            }
        )
    )

    assert should_parse(source_path, meta_path) is True


def test_should_parse_retries_current_cached_failure(tmp_path):
    source_path = tmp_path / "failed.pdf"
    source_path.write_bytes(b"%PDF-1.4 failed")
    meta_path = tmp_path / "failed.meta.json"
    meta_path.write_text(
        json.dumps(
            {
                "source_file_hash": f"sha256:{__import__('hashlib').sha256(source_path.read_bytes()).hexdigest()}",
                "parse_cache_version": PARSE_CACHE_VERSION,
                "parse_status": "parse_exception",
            }
        )
    )

    assert should_parse(source_path, meta_path) is True


def test_parse_failure_sidecar_is_versioned_typed_and_counts_attempts(
    monkeypatch, tmp_path
):
    source_path = tmp_path / "contract.pdf"
    source_path.write_bytes(b"%PDF-1.4 fake")
    meta_path = tmp_path / "contract.meta.json"
    source_hash = (
        f"sha256:{__import__('hashlib').sha256(source_path.read_bytes()).hexdigest()}"
    )
    meta_path.write_text(
        json.dumps(
            {
                "attempt_count": 2,
                "source_file_hash": source_hash,
                "parse_cache_version": PARSE_CACHE_VERSION,
            }
        )
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda *_args: (
            tmp_path / "contract.json",
            tmp_path / "contract.md",
            meta_path,
        ),
    )
    from pipeline.parse.converter import save_parse_error

    save_parse_error(source_path, "managed_care_contract", TimeoutError("slow"))
    meta = json.loads(meta_path.read_text())

    assert meta["parse_cache_version"] == PARSE_CACHE_VERSION
    assert meta["failure_kind"] == "timeout"
    assert meta["attempt_count"] == 3


def test_inventory_rejects_cached_failure_even_when_current_run_skips_it(
    monkeypatch, tmp_path
):
    source = tmp_path / "failed.pdf"
    source.write_bytes(b"pdf")
    meta = tmp_path / "failed.meta.json"
    meta.write_text(json.dumps({"parse_status": "parse_exception"}))
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda *_args: (tmp_path / "failed.json", tmp_path / "failed.md", meta),
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )

    failures = inventory_parse_failures([source])

    assert failures[0]["source_file"] == str(source)
    assert failures[0]["parse_status"] == "parse_exception"
    assert failures[0]["required_for_answering"] is True


def test_inventory_excludes_removed_and_orphan_registry_files(monkeypatch, tmp_path):
    active = tmp_path / "active.pdf"
    removed = tmp_path / "removed.pdf"
    orphan = tmp_path / "orphan.pdf"
    for source in (active, removed, orphan):
        source.write_bytes(b"pdf")
    meta_paths = {
        source: tmp_path / f"{source.stem}.meta.json"
        for source in (active, removed, orphan)
    }
    for meta in meta_paths.values():
        meta.write_text(json.dumps({"parse_status": "parse_exception"}))
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda source, _type: (
            tmp_path / f"{source.stem}.json",
            tmp_path / f"{source.stem}.md",
            meta_paths[source],
        ),
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.read_registry_rows",
        lambda: [
            {"local_path": str(active), "status": "unchanged"},
            {"local_path": str(removed), "status": "removed"},
        ],
    )

    failures = inventory_parse_failures([active, removed, orphan])

    assert [failure["source_file"] for failure in failures] == [str(active)]


def test_inventory_blocks_active_registry_source_missing_from_disk(
    monkeypatch, tmp_path
):
    missing = tmp_path / "clinical-policies" / "missing.pdf"
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing-registry.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.read_registry_rows",
        lambda: [{"local_path": str(missing), "status": "downloaded"}],
    )

    failures = inventory_parse_failures([])
    evidence = build_parse_inventory_evidence([])

    assert failures == [
        {
            "source_file": str(missing),
            "source_type": "clinical_policy",
            "parse_status": "missing_source_file",
            "required_for_answering": True,
            "attempt_count_generation": 0,
            "retryable": False,
        }
    ]
    assert evidence["active_source_count"] == 1
    assert evidence["unresolved_required_count"] == 1
    assert evidence["unresolved"][0]["source_hash"] is None


def test_inventory_keeps_retryable_crawler_error_rows_active(monkeypatch, tmp_path):
    source = tmp_path / "clinical-policies" / "errored.pdf"
    source.parent.mkdir()
    source.write_bytes(b"%PDF-1.4 retained source")
    outputs = (
        tmp_path / "errored.json",
        tmp_path / "errored.md",
        tmp_path / "errored.meta.json",
    )
    outputs[2].write_text(
        json.dumps(
            {
                "parse_status": "parse_exception",
                "attempt_count_generation": 2,
            }
        )
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "registry.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.read_registry_rows",
        lambda: [{"local_path": str(source), "status": "error"}],
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths", lambda *_args: outputs
    )

    failures = inventory_parse_failures([])
    evidence = build_parse_inventory_evidence([])

    assert failures[0]["source_file"] == str(source)
    assert failures[0]["parse_status"] == "parse_exception"
    assert failures[0]["required_for_answering"] is True
    assert evidence["active_source_count"] == 1
    assert evidence["unresolved_required_count"] == 1


def test_inventory_fails_closed_for_existing_empty_registry(monkeypatch, tmp_path):
    registry = tmp_path / "crawler.db"
    registry.write_bytes(b"")
    monkeypatch.setattr("pipeline.parse.converter.CRAWLER_DB_PATH", registry)
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])

    failures = inventory_parse_failures([])

    assert failures == [
        {
            "source_file": str(registry),
            "source_type": "crawler_registry",
            "parse_status": "empty_registry",
            "required_for_answering": True,
            "attempt_count_generation": 0,
            "retryable": False,
        }
    ]


def test_inventory_hash_binds_artifact_and_metadata_bytes(monkeypatch, tmp_path):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"workbook")
    outputs = (
        tmp_path / "source.json",
        tmp_path / "source.md",
        tmp_path / "source.meta.json",
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file", lambda _path: "test"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths", lambda *_args: outputs
    )

    first = _parsed_fixture()
    save_parsed(source, first, "test")
    first_evidence = build_parse_inventory_evidence([source])

    second = _parsed_fixture()
    second["markdown"] = "different substantive fixture markdown"
    second["payload"] = {"name": "different"}
    save_parsed(source, second, "test")
    second_evidence = build_parse_inventory_evidence([source])

    assert first_evidence["inventory_hash"] != second_evidence["inventory_hash"]
    assert first_evidence["unresolved_count"] == 0
    assert second_evidence["unresolved_count"] == 0


def test_pdf_uses_no_ocr_first_and_falls_back_for_blank_output(monkeypatch, tmp_path):
    source_path = tmp_path / "scan.pdf"
    source_path.write_bytes(b"pdf")

    class BlankDoc(_FakeDoc):
        def export_to_dict(self):
            return {"tables": []}

        def export_to_markdown(self):
            return ""

        def iterate_items(self):
            return iter(())

    class TextDoc(BlankDoc):
        def export_to_markdown(self):
            return "Recovered text"

        def iterate_items(self):
            return iter([(_FakeDocItem(DocItemLabel.TEXT, "Recovered text"), 0)])

    class Converter:
        def __init__(self, doc):
            self.doc = doc

        def convert(self, _path):
            result = _FakeConversionResult()
            result.document = self.doc
            return result

    created = []

    def fake_create_converter(*, do_ocr=False):
        created.append(do_ocr)
        return Converter(TextDoc())

    monkeypatch.setattr(
        "pipeline.parse.converter.create_converter", fake_create_converter
    )
    parsed = parse_pdf(source_path, Converter(BlankDoc()))

    assert created == [True]
    assert parsed["ocr_attempted"] is True
    assert parsed["ocr_used"] is True
    assert [attempt["ocr"] for attempt in parsed["parse_attempts"]] == [False, True]
    assert parsed["answer_ready"] is True


def test_create_converter_defaults_to_no_ocr(monkeypatch):
    captured = {}

    class FakeDocumentConverter:
        def __init__(self, *, format_options):
            captured.update(format_options)

    monkeypatch.setattr(
        "pipeline.parse.converter.DocumentConverter", FakeDocumentConverter
    )
    create_converter()
    options = next(iter(captured.values())).pipeline_options
    assert options.do_ocr is False


def test_parse_all_records_timeout_failures(monkeypatch, tmp_path):
    source_path = tmp_path / "contract.pdf"
    source_path.write_bytes(b"%PDF-1.4 fake")
    meta_path = tmp_path / "contract.meta.json"
    monkeypatch.setattr(
        "pipeline.parse.converter.PIPELINE_DATA_DIR", tmp_path / "pipeline"
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )

    monkeypatch.setattr(
        "pipeline.parse.converter.discover_source_files",
        lambda _source_type=None: [source_path],
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file",
        lambda _source_path: "managed_care_contract",
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda _source_path, _source_type: (
            tmp_path / "contract.json",
            tmp_path / "contract.md",
            meta_path,
        ),
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parse_and_save_supported_file_with_timeout",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            TimeoutError("parse timed out after 1 seconds")
        ),
    )

    stats = parse_all(source_type="managed_care_contract")
    meta = json.loads(meta_path.read_text())

    assert {
        key: value for key, value in stats.items() if key != "inventory_evidence"
    } == {
        "parsed": 0,
        "skipped": 0,
        "errors": 1,
        "timed_out": 1,
        "quarantined": 0,
        "remaining_failures": 1,
        "remaining_required_failures": 0,
        "remaining_optional_failures": 1,
        "terminal_state": "completed",
    }
    assert meta["parse_status"] == "parse_exception"
    assert meta["answer_ready"] is False
    assert "timed out after 1 seconds" in meta["error_message"]


def test_parse_all_quarantines_unsupported_formats(monkeypatch, tmp_path):
    source_path = tmp_path / "legacy-slides.ppt"
    source_path.write_bytes(b"legacy presentation")
    meta_path = tmp_path / "legacy-slides__ppt.meta.json"
    monkeypatch.setattr(
        "pipeline.parse.converter.PIPELINE_DATA_DIR", tmp_path / "pipeline"
    )

    monkeypatch.setattr(
        "pipeline.parse.converter.discover_source_files",
        lambda _source_type=None: [],
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_unsupported_files",
        lambda _source_type=None: [source_path],
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file",
        lambda _source_path: "managed_care_guidance",
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda _source_path, _source_type: (
            tmp_path / "legacy-slides__ppt.json",
            tmp_path / "legacy-slides__ppt.md",
            meta_path,
        ),
    )

    stats = parse_all(source_type="managed_care_guidance")
    meta = json.loads(meta_path.read_text())

    assert stats["quarantined"] == 1
    assert stats["errors"] == 0
    assert meta["parse_status"] == "unsupported_format"
    assert meta["answer_ready"] is False
    assert "artifact quarantined" in meta["error_message"]


def test_parse_all_retry_failures_limit_and_atomic_terminal_state(
    monkeypatch, tmp_path
):
    sources = [tmp_path / name for name in ("ok.xlsx", "failed.xlsx", "failed-2.xlsx")]
    for source in sources:
        source.write_bytes(b"workbook")
    meta_paths = {source: tmp_path / f"{source.stem}.meta.json" for source in sources}
    meta_paths[sources[0]].write_text(
        json.dumps({"parse_status": "parsed", "answer_ready": True})
    )
    for source in sources[1:]:
        meta_paths[source].write_text(json.dumps({"parse_status": "parse_exception"}))

    monkeypatch.setattr(
        "pipeline.parse.converter.PIPELINE_DATA_DIR", tmp_path / "pipeline"
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_source_files", lambda _source=None: sources
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_unsupported_files", lambda _source=None: []
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file", lambda _path: "test"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda source, _type: (
            tmp_path / "out.json",
            tmp_path / "out.md",
            meta_paths[source],
        ),
    )
    parsed = []

    def fake_parse(source, _source_type, _converter):
        parsed.append(source)
        meta_paths[source].write_text(
            json.dumps({"parse_status": "parsed", "answer_ready": True})
        )
        return {"section_count": 1, "table_count": 0, "parse_time_seconds": 0}

    monkeypatch.setattr(
        "pipeline.parse.converter.parse_and_save_supported_file", fake_parse
    )

    stats = parse_all(retry_failures=True, limit=1)
    state_path = tmp_path / "pipeline" / "run-state" / "latest-parse.json"
    state = json.loads(state_path.read_text())

    assert parsed == [sources[2]]
    assert stats["remaining_failures"] == 3
    assert stats["terminal_state"] == "completed"
    assert state["status"] == "completed"
    failures = {
        Path(item["source_file"]): item["parse_status"]
        for item in state["remaining_failures"]
    }
    assert failures == {
        sources[0]: "stale_sidecar",
        sources[1]: "parse_exception",
        sources[2]: "stale_sidecar",
    }
    assert not state_path.with_suffix(".tmp").exists()


def test_parse_all_retry_file_selector_does_not_expand_to_active_registry(
    monkeypatch, tmp_path
):
    selected = tmp_path / "clinical-policies" / "selected.pdf"
    unrelated = tmp_path / "clinical-policies" / "unrelated.pdf"
    selected.parent.mkdir()
    selected.write_bytes(b"selected")
    unrelated.write_bytes(b"unrelated")
    metadata = {
        selected: tmp_path / "selected.meta.json",
        unrelated: tmp_path / "unrelated.meta.json",
    }
    for path in metadata.values():
        path.write_text(json.dumps({"parse_status": "parse_exception"}))

    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "registry.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.read_registry_rows",
        lambda: [
            {"local_path": str(selected), "status": "downloaded"},
            {"local_path": str(unrelated), "status": "downloaded"},
        ],
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda source, *_args: (
            tmp_path / f"{source.stem}.json",
            tmp_path / f"{source.stem}.md",
            metadata[source],
        ),
    )
    monkeypatch.setattr(
        "pipeline.parse.converter._write_run_state", lambda _state: None
    )
    parsed = []
    monkeypatch.setattr(
        "pipeline.parse.converter.parse_and_save_supported_file_with_timeout",
        lambda source, *_args: (
            parsed.append(source)
            or {"section_count": 1, "table_count": 0, "parse_time_seconds": 0.1}
        ),
    )

    stats = parse_all(file_path=selected, retry_failures=True)

    assert parsed == [selected]
    assert stats["remaining_failures"] == 1
    assert stats["inventory_evidence"]["active_source_count"] == 1


def test_retry_source_selector_inventories_only_selected_active_family(
    monkeypatch, tmp_path
):
    clinical = tmp_path / "clinical-policies" / "missing.pdf"
    state_plan = tmp_path / "state-plan" / "missing.pdf"
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "registry.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.read_registry_rows",
        lambda: [
            {"local_path": str(clinical), "status": "downloaded"},
            {"local_path": str(state_plan), "status": "downloaded"},
        ],
    )

    failures = inventory_parse_failures([], selected_source_type="clinical_policy")

    assert [failure["source_file"] for failure in failures] == [str(clinical)]
    assert failures[0]["parse_status"] == "missing_source_file"


def test_parse_all_stops_after_consecutive_timeouts(monkeypatch, tmp_path):
    sources = [tmp_path / f"{index}.pdf" for index in range(3)]
    for source in sources:
        source.write_bytes(b"pdf")
    monkeypatch.setattr(
        "pipeline.parse.converter.PIPELINE_DATA_DIR", tmp_path / "pipeline"
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_source_files", lambda _source=None: sources
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_unsupported_files", lambda _source=None: []
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file", lambda _path: "test"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths",
        lambda source, _type: (
            tmp_path / "out.json",
            tmp_path / "out.md",
            tmp_path / f"{source.stem}.meta.json",
        ),
    )
    calls = []

    def timeout(source, *_args):
        calls.append(source)
        raise TimeoutError("slow")

    monkeypatch.setattr(
        "pipeline.parse.converter.parse_and_save_supported_file_with_timeout", timeout
    )

    stats = parse_all(max_consecutive_timeouts=2)

    assert calls == sources[:2]
    assert stats["terminal_state"] == "max_consecutive_timeouts_reached"


def test_parse_all_preserves_valid_commit_when_worker_result_is_lost(
    monkeypatch, tmp_path
):
    source = tmp_path / "committed.xlsx"
    source.write_bytes(b"workbook")
    outputs = (
        tmp_path / "committed.json",
        tmp_path / "committed.md",
        tmp_path / "committed.meta.json",
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.PIPELINE_DATA_DIR", tmp_path / "pipeline"
    )
    monkeypatch.setattr("pipeline.parse.converter.read_registry_rows", lambda: [])
    monkeypatch.setattr(
        "pipeline.parse.converter.CRAWLER_DB_PATH", tmp_path / "missing.db"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_source_files", lambda _source=None: [source]
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.discover_unsupported_files", lambda _source=None: []
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.source_type_for_file", lambda _path: "test"
    )
    monkeypatch.setattr(
        "pipeline.parse.converter.parsed_output_paths", lambda *_args: outputs
    )

    def commit_then_lose_result(source_path, source_type, _converter):
        save_parsed(source_path, _parsed_fixture(), source_type)
        raise RuntimeError("worker result lost")

    monkeypatch.setattr(
        "pipeline.parse.converter.parse_and_save_supported_file",
        commit_then_lose_result,
    )

    stats = parse_all()
    meta = json.loads(outputs[2].read_text())

    assert stats["parsed"] == 1
    assert stats["errors"] == 0
    assert stats["remaining_failures"] == 0
    assert meta["attempt_count_total"] == 1
